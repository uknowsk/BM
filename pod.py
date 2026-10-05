"""Normalize free-text POD (selling point) features into a fixed comparison taxonomy.

Strategy: keyword rules first, one batched LLM call per product for leftovers,
answers cached by sha1(raw_text) in pod_cache.json.
"""
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from typing import Callable, Literal, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

import catalog
from excel_writer import clean_text, is_formula_like
from schema import ProductRecord

CACHE_PATH = Path(__file__).resolve().parent / "pod_cache.json"
OTHER = "other"

DEFAULT_MAJOR = "refrigerator"

# (key, category, English name, Korean name) -- order defines matrix order.
# One taxonomy per major product group (catalog.CATEGORY_TREE keys); TAXONOMY is the refrigerator one.
TAXONOMY: list[tuple[str, str, str, str]] = [
    ("fresh_tech", "Freshness", "Fresh-food preservation technology", "신선 보존 기술"),
    ("humidity_drawer", "Freshness", "Humidity-controlled crisper drawer", "습도 조절 야채실"),
    ("airflow_temp", "Freshness", "Airflow / temperature management", "냉기 순환·온도 관리"),
    ("freshness_mode", "Freshness", "Freshness mode / setting", "프레시니스 모드"),
    ("dual_cooling", "Cooling System", "Dual compressor / dual evaporator", "듀얼 컴프레서·증발기"),
    ("super_cool_freeze", "Cooling System", "Super cool / turbo cool / fast freeze", "급속 냉각·냉동"),
    ("no_frost", "Cooling System", "No-frost / auto defrost", "성에 방지(무서리)"),
    ("ice_maker", "Ice & Water", "Ice maker", "제빙기"),
    ("water_dispenser", "Ice & Water", "Water dispenser", "정수(물) 디스펜서"),
    ("water_filter", "Ice & Water", "Water filter", "정수 필터"),
    ("ice_type", "Ice & Water", "Crushed / cubed ice options", "얼음 종류 선택(조각/크러시)"),
    ("wifi_app", "Smart", "Wi-Fi / app remote control", "Wi-Fi·앱 원격 제어"),
    ("smart_notify", "Smart", "Smart notifications / voice assistant", "스마트 알림·음성 지원"),
    ("door_alarm", "Convenience", "Door-ajar alarm", "도어 열림 알림"),
    ("child_lock", "Convenience", "Child lock / control lock", "잠금(어린이 보호)"),
    ("temp_display", "Convenience", "Temperature display / external controls", "온도 표시·외부 조작부"),
    ("sabbath_mode", "Convenience", "Sabbath mode", "안식일 모드"),
    ("vacation_mode", "Convenience", "Vacation / away mode", "휴가 모드"),
    ("filter_reminder", "Convenience", "Filter change reminder", "필터 교체 알림"),
    ("flex_space", "Convenience", "Quick-space / flex shelf or zone", "퀵스페이스·가변 공간"),
    ("door_bins", "Storage", "Gallon / adjustable door bins", "갤런 도어 바스켓"),
    ("shelves", "Storage", "Spill-proof / adjustable shelves", "스필 방지·조절 선반"),
    ("freezer_storage", "Storage", "Freezer baskets / drawers", "냉동실 바스켓·서랍"),
    ("finish", "Design", "Fingerprint-resistant / stainless finish", "지문 방지·스테인리스 마감"),
    ("led_light", "Design", "LED lighting", "LED 조명"),
    ("steel_back", "Design", "Stainless back wall", "스테인리스 내벽"),
    ("energy_star", "Energy", "ENERGY STAR", "에너지스타"),
    ("eco_energy_saving_mode", "Energy", "Eco / energy-saving mode", "절전(에너지 절약) 모드"),
    (OTHER, "Other", "Other", "기타"),
]

WASHER_TAXONOMY: list[tuple[str, str, str, str]] = [
    ("steam", "Washing", "Steam wash / steam care", "스팀 세탁·스팀 케어"),
    ("sanitize", "Washing", "Sanitize cycle / bacteria removal", "살균·위생 코스"),
    ("allergen", "Washing", "Allergen cycle", "알러지 케어 코스"),
    ("drum_clean", "Washing", "Tub / drum self-clean", "통세척·자동 세척"),
    ("quick_wash", "Washing", "Quick / speed wash", "쾌속·빠른 세탁 코스"),
    ("bulky", "Washing", "Bulky / bedding / heavy-duty cycle", "이불·대용량 코스"),
    ("extra_rinse", "Washing", "Extra rinse / deep fill", "추가 헹굼·딥필"),
    ("auto_dispense", "Convenience", "Auto-dispense detergent", "자동 세제 투입"),
    ("delay_cycle", "Convenience", "Delay start / end", "예약 세탁"),
    ("child_lock", "Convenience", "Child / control lock", "어린이 보호 잠금"),
    ("display_controls", "Convenience", "Display / touch controls", "디스플레이·조작부"),
    ("cycles", "Convenience", "Cycle options / custom programs", "세탁 코스·옵션"),
    ("wifi_app", "Smart", "Wi-Fi / app remote control", "Wi-Fi·앱 원격 제어"),
    ("voice_assist", "Smart", "Voice assistant", "음성 비서 지원"),
    ("ai_sensing", "Smart", "AI / load & fabric sensing", "AI·세탁물 자동 감지"),
    ("stainless_drum", "Drum & Design", "Stainless steel drum", "스테인리스 세탁조"),
    ("large_capacity", "Drum & Design", "Large capacity", "대용량"),
    ("stackable", "Drum & Design", "Stackable / space-saving", "스택·공간 절약"),
    ("finish", "Drum & Design", "Finish / color", "마감·색상"),
    ("dryer_sensor", "Drying", "Moisture / sensor dry", "습도 센서 건조"),
    ("heat_pump", "Drying", "Heat pump / condensing dryer", "히트펌프·콘덴싱 건조"),
    ("wrinkle_care", "Drying", "Wrinkle prevention / refresh", "구김 방지·리프레시"),
    ("vibration_noise", "Noise", "Vibration / noise reduction", "진동·소음 저감"),
    ("energy_star", "Energy", "ENERGY STAR", "에너지스타"),
    ("eco_saving", "Energy", "Eco / water & energy saving", "절수·절전"),
    (OTHER, "Other", "Other", "기타"),
]

COOKING_TAXONOMY: list[tuple[str, str, str, str]] = [
    ("convection", "Cooking Tech", "Convection", "컨벡션"),
    ("air_fry", "Cooking Tech", "Air fry", "에어프라이"),
    ("steam_cook", "Cooking Tech", "Steam cooking / sous vide", "스팀 조리"),
    ("fast_preheat", "Cooking Tech", "Fast / no preheat", "빠른 예열"),
    ("temp_probe", "Cooking Tech", "Temperature probe", "온도 프로브"),
    ("power_levels", "Cooking Tech", "Power levels / inverter (microwave)", "출력 단계·인버터"),
    ("sensor_cook", "Cooking Tech", "Sensor / auto cook", "센서·자동 조리"),
    ("microwave_power", "Cooking Tech", "Microwave power", "마이크로웨이브 출력"),
    ("speed_cook", "Cooking Tech", "Speed cook / combi modes (microwave + convection / light wave)", "스피드쿡·복합 조리 모드"),
    ("induction_tech", "Cooktop", "Induction heating / pan detection", "인덕션 가열·용기 감지"),
    ("smoothtop", "Cooktop", "Smooth glass-ceramic cooktop", "라디언트·글래스 상판"),
    ("gas_burner", "Cooktop", "Sealed / auto-reignite gas burners", "가스 버너·자동 재점화"),
    ("boost_burner", "Cooktop", "Boost / power burner", "부스트·고화력 버너"),
    ("flex_zone", "Cooktop", "Bridge / flex cooking zone", "브릿지·가변 쿠킹존"),
    ("griddle", "Cooktop", "Griddle / center oval burner", "그리들·센터 오벌"),
    ("continuous_grates", "Cooktop", "Continuous grates", "연속 그레이트"),
    ("double_oven", "Oven Layout", "Double / dual-cavity oven", "더블·듀얼 오븐"),
    ("warming_drawer", "Oven Layout", "Warming / storage drawer", "워밍·수납 서랍"),
    ("capacity", "Oven Layout", "Large cavity / capacity", "대용량 캐비티"),
    ("self_clean", "Cleaning", "Self-clean / steam clean", "자동 세척(셀프 클린)"),
    ("wifi_app", "Smart", "Wi-Fi / app remote control", "Wi-Fi·앱 원격 제어"),
    ("voice_assist", "Smart", "Voice assistant", "음성 비서 지원"),
    ("sabbath_mode", "Convenience", "Sabbath mode", "안식일 모드"),
    ("child_lock", "Convenience", "Child / control lock", "어린이 보호 잠금"),
    ("timer_safety", "Convenience", "Timer / auto shut-off", "타이머·자동 차단"),
    ("controls", "Convenience", "Display / knob controls", "디스플레이·조작부"),
    ("ventilation", "Ventilation & Light", "Ventilation / exhaust fan", "배기·환기 팬"),
    ("cooktop_light", "Ventilation & Light", "Cooktop / cavity lighting", "조명"),
    ("finish", "Design", "Finish / design", "마감·디자인"),
    ("energy_star", "Energy", "ENERGY STAR", "에너지스타"),
    (OTHER, "Other", "Other", "기타"),
]

TAXONOMIES = {"refrigerator": TAXONOMY, "washer": WASHER_TAXONOMY, "cooking": COOKING_TAXONOMY}
_TAXES = {major: {t[0]: t for t in tax} for major, tax in TAXONOMIES.items()}
_TAX = _TAXES[DEFAULT_MAJOR]


def major_of_product(product: ProductRecord) -> str:
    """Major product group of a record: ProductRecord.category (key or legacy label), else the major of its
    subcategory key, else the refrigerator default."""
    return (catalog.normalize_major(product.category) or catalog.major_of(product.subcategory or "")
            or DEFAULT_MAJOR)

# Every rule that matches contributes a key (bundles yield several). Rules are
# (key, regex, weak): weak rules only count when no strong rule matched.
_RULES: list[tuple[str, str, bool]] = [
    ("sabbath_mode", r"sabbath|shabbat|shabbos", False),
    ("energy_star", r"energy\s*star", False),
    ("filter_reminder", r"filter (change|replace|status|reminder|notification|indicator)|(change|replace)[- ]filter", False),
    ("super_cool_freeze", r"super[- ]?(cool|freez)|turbo[- ]?cool|fast[- ]?free|power[- ]?(cool|freez)|quick[- ]?(freez|chill)|rapid(ly)?[- ]?(cool|freez)", False),
    ("vacation_mode", r"vacation|holiday|away mode", False),
    ("eco_energy_saving_mode", r"energy[- ]?sav|power[- ]?sav|\beco\b", False),
    ("freshness_mode", r"fresh(ness)? (mode|setting)", False),
    ("wifi_app", r"wi-?fi|smartthings|home ?connect|\bapp\b|remote (control|monitor)", False),
    ("door_alarm", r"door[- ]?(ajar|alarm|open)|open door|alarm", False),
    ("smart_notify", r"alexa|google assistant|voice|notification|smart (home|alert)", False),
    ("child_lock", r"child|\block\b|locking|lockout", False),
    ("fresh_tech", r"vitafresh|freshlock|fresh ?(food )?(preserv|keep|tech)|keeps? (food|produce) fresh|preserv|farm ?fresh|extended freshness|ethylene", False),
    ("humidity_drawer", r"humidity|crisper|produce drawer|vegetable drawer|moisture", False),
    ("airflow_temp", r"airflow|air flow|air circulation|air tower|temperature (management|control|zone)|precise temp|multi-?flow|temp(erature)?-?controlled", False),
    ("dual_cooling", r"dual[- ]?(compressor|evaporator|cool)|twin cooling|two evaporator|separate evaporator", False),
    ("no_frost", r"no[- ]?frost|frost[- ]?free|auto(matic)? defrost|defrost", False),
    ("ice_type", r"crushed|cubed|ice (type|option|selection)|cube ice|nugget", False),
    ("ice_maker", r"ice ?maker|ice ?machine|automatic ice", False),
    ("ice_maker", r"\bice\b", True),
    ("water_dispenser", r"water dispens|dispenser", False),
    ("water_filter", r"(?<!ethylene )filter|filtrat|bypass (cap|plug)", False),
    ("water_dispenser", r"\bwater\b", True),
    ("temp_display", r"display|digital (control|temp)|external (control|temp)|touch ?(screen|control)|control panel|temperature (indicator|readout)", False),
    ("flex_space", r"quick[- ]?space|flex|convertible|variable|custom(izable)? (zone|drawer|space)|beverage center|snack drawer", False),
    ("door_bins", r"gallon|door (bin|storage|basket)|bins?\b", False),
    ("freezer_storage", r"freezer (basket|drawer|storage|shelf)|basket|slide[- ]?out", False),
    ("shelves", r"shelf|shelves|spill|glass shelf|adjustable", False),
    ("steel_back", r"stainless (steel )?(back|interior|inner|wall)|steel (back|interior)", False),
    ("finish", r"fingerprint|smudge|stainless|finish|slate|black steel", False),
    ("led_light", r"\bled\b|lighting", False),
    ("led_light", r"light(s)?\b", True),
]
_WASHER_RULES: list[tuple[str, str, bool]] = [
    ("energy_star", r"energy\s*star", False),
    ("sanitize", r"saniti[sz]|sterili[sz]|bacteria|germ|antibacterial", False),
    ("allergen", r"allergen|allergy|allergies", False),
    ("steam", r"steam", False),
    ("drum_clean", r"(tub|drum) clean|self[- ]?clean|clean(s|ing)? (the )?(tub|drum)|washer clean|affresh", False),
    ("wifi_app", r"wi-?fi|smartthings|thinq|home ?connect|\bapp\b|remote (start|control|monitor)|smart pairing", False),
    ("voice_assist", r"alexa|google assistant|bixby|voice", False),
    ("auto_dispense", r"auto[- ]?(dispens|dos)|automatic(ally)? (detergent|dispens|dos)|detergent (dispens|dosing)|smart dispens|ez ?dispense|load ?& ?go", False),
    ("stainless_drum", r"stainless[- ]?(steel )?(drum|tub|basket|wash basket)|steel (drum|tub|basket)", False),
    ("quick_wash", r"quick|speed[- ]?(wash|cycle)|\d+[- ]?min(ute)?s? (cycle|wash)|fast (wash|cycle)|turbo ?wash|super ?speed", False),
    ("vibration_noise", r"vibration|quiet|noise|sound (reduc|insul)|stabiliz", False),
    ("stackable", r"stack", False),
    ("ai_sensing", r"\bai\b|load[- ]?sens|sensing|auto[- ]?(load|adjust)|fabric sens|detects?\b", False),
    ("large_capacity", r"capacity|cu\.? ?ft", False),
    ("heat_pump", r"heat[- ]?pump|dual inverter|condens", False),
    ("dryer_sensor", r"moisture sens|sensor dry|dryness|auto(matic)? dry", False),
    ("dryer_sensor", r"sensor", True),
    ("wrinkle_care", r"wrinkle|refresh|crease|tumble press", False),
    ("delay_cycle", r"delay (start|wash|end)|time delay|schedul", False),
    ("child_lock", r"child|control lock|\block\b|lockout", False),
    ("bulky", r"bulky|comforter|bedding|heavy[- ]?duty", False),
    ("extra_rinse", r"extra rinse|deep (fill|water)|rinse", False),
    ("eco_saving", r"water[- ]?(sav|effic)|high[- ]efficiency|\bhe\b|eco ?(bubble|hybrid)|\beco\b|energy[- ]?sav", False),
    ("display_controls", r"display|touch|\blcd\b|\bled\b|knob|dial|control panel", False),
    ("finish", r"finish|fingerprint|smudge|color|graphite|stainless", False),
    ("cycles", r"cycle|option|program|setting|course", True),
]
_COOKING_RULES: list[tuple[str, str, bool]] = [
    ("sabbath_mode", r"sabbath|shabbat|shabbos", False),
    ("energy_star", r"energy\s*star", False),
    ("self_clean", r"self[- ]?clean|steam clean|easy ?clean|pyrolytic|auto ?clean|clean cycle", False),
    ("convection", r"convection|fan[- ]?assist", False),
    ("air_fry", r"air[- ]?fr(y|ier)|airfry|crisp", False),
    ("steam_cook", r"steam|sous ?vide", False),
    ("fast_preheat", r"preheat|pre-heat", False),
    ("temp_probe", r"probe|thermometer", False),
    ("sensor_cook", r"sensor ?(cook|reheat)|auto(matic)? (cook|defrost|reheat)|one[- ]touch|smart cook|sensor", False),
    ("speed_cook", r"speed ?cook|speed oven|combi(nation)?|microwave (and|\+|with) convection|microwave convection|light[- ]?wave|halogen|qooker|turbo ?cook", False),
    ("microwave_power", r"microwave (power|output|wattage)|\d{3,4} ?watt microwave|microwave \d{3,4} ?w", False),
    ("power_levels", r"power (level|setting)|inverter|\d{3,4} ?watt|\d{3,4}w\b|wattage", False),
    ("wifi_app", r"wi-?fi|smartthings|thinq|home ?connect|\bapp\b|remote|connected", False),
    ("voice_assist", r"alexa|google assistant|bixby|voice", False),
    ("induction_tech", r"induction|pan (detect|sens)|cookware sens|magnetic", False),
    ("smoothtop", r"radiant|ceramic|smooth ?top|glass[- ]?top|glass cooktop", False),
    ("gas_burner", r"re-?ignit|spark|ignit|sealed burner|gas burner|flame", False),
    ("boost_burner", r"boost|power ?boil|rapid|\bbtu\b|tri[- ]?ring|triple|dual[- ]?(ring|stacked|flame)|power burner|high[- ]?power", False),
    ("flex_zone", r"bridge|flex|expandable|multi[- ]?element|dual element|triple element|two[- ]?in[- ]?one element|cooking zone", False),
    ("griddle", r"griddle|center oval|oval burner|reversible grill|grill", False),
    ("continuous_grates", r"grate|edge[- ]to[- ]edge|continuous", False),
    ("double_oven", r"double oven|dual oven|second oven|two ovens|upper oven|lower oven|dual cavity|oven divider", False),
    ("warming_drawer", r"warming|storage drawer|proof", False),
    ("capacity", r"capacity|cu\.? ?ft|large (oven|cavity)|extra[- ]?large", False),
    ("child_lock", r"child|control lock|lockout|\block\b", False),
    ("timer_safety", r"timer|keep warm|shut-?off|auto[- ]?off", False),
    ("ventilation", r"\bvent|exhaust|(?<!with )(?<!assist )\bfan\b|\bcfm\b|\bduct|recirculat|\bhood\b", False),
    ("cooktop_light", r"\bled\b|light(ing|s)?\b", False),
    ("controls", r"display|touch|knob|control panel|digital|\bdial\b", False),
    ("finish", r"stainless|finish|fingerprint|smudge|slate|black|design|handle", False),
]
_COMPILED = [(k, re.compile(p, re.I), w) for k, p, w in _RULES]
_COMPILED_BY = {
    "refrigerator": _COMPILED,
    "washer": [(k, re.compile(p, re.I), w) for k, p, w in _WASHER_RULES],
    "cooking": [(k, re.compile(p, re.I), w) for k, p, w in _COOKING_RULES],
}
_SUPPRESS_EXTRA = {
    "washer": {
        "stainless_drum": {"finish"}, "quick_wash": {"cycles"}, "stackable": {"large_capacity"},
        "bulky": {"extra_rinse", "large_capacity"}, "heat_pump": {"dryer_sensor"},
    },
    "cooking": {
        "self_clean": {"steam_cook"}, "double_oven": {"capacity"}, "induction_tech": {"smoothtop"},
        "sensor_cook": {"controls"},
    },
}
# key -> keys it makes redundant when both match the same text
_SUPPRESS = {
    "flex_space": {"shelves"},
    "ice_type": {"ice_maker"},
    "steel_back": {"finish"},
    "door_alarm": {"smart_notify"},
    "filter_reminder": {"smart_notify", "water_filter"},
    "led_light": {"water_dispenser", "ice_maker"},
}
_SUPPRESS_BY = {"refrigerator": _SUPPRESS, **_SUPPRESS_EXTRA}

_FIELD_EVIDENCE = (  # (ProductRecord attr, taxonomy key, synthetic raw text)
    ("wifi_supported", "wifi_app", "Wi-Fi supported (spec field)"),
    ("ice_maker", "ice_maker", "Ice maker (spec field)"),
    ("water_dispenser", "water_dispenser", "Water dispenser (spec field)"),
    ("energy_star", "energy_star", "ENERGY STAR (spec field)"),
)
_FIELD_EVIDENCE_BY = {
    "refrigerator": _FIELD_EVIDENCE,
    "washer": (_FIELD_EVIDENCE[0], _FIELD_EVIDENCE[3]),
    "cooking": (_FIELD_EVIDENCE[0], _FIELD_EVIDENCE[3]),
}


class PodItem(BaseModel):
    brand: str
    model_number: str
    raw_text: str
    taxonomy_key: str
    category: str
    name_en: str
    name_ko: str
    matched_by: Literal["keyword", "llm", "none", "mode"]
    mode_ref: Optional[str] = None  # e.g. 'Modes: Sabbath Mode p12' (mode evidence for this key)


def keyword_keys(text: str, category: str = DEFAULT_MAJOR) -> list[str]:
    """All taxonomy keys (of the major group's taxonomy) whose rules match text (rule order, de-duplicated)."""
    text = (text or "").strip()
    rules, suppress = _COMPILED_BY[category], _SUPPRESS_BY[category]
    strong = [k for k, rx, weak in rules if not weak and rx.search(text)]
    found = strong or [k for k, rx, weak in rules if weak and rx.search(text)]
    found = list(dict.fromkeys(found))
    drop = set().union(*(suppress.get(k, set()) for k in found)) if found else set()
    return [k for k in found if k not in drop]


def keyword_key(text: str) -> Optional[str]:
    """First matching key (kept for backward compatibility)."""
    keys = keyword_keys(text)
    return keys[0] if keys else None


def _sha(text: str, category: str = DEFAULT_MAJOR) -> str:
    """Cache key; refrigerator keys stay unprefixed so existing pod_cache.json entries remain valid."""
    return hashlib.sha1((text if category == DEFAULT_MAJOR else f"{category}:{text}").encode("utf-8")).hexdigest()


def _load_cache(path: Path) -> dict[str, str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _save_cache(path: Path, cache: dict[str, str]) -> None:
    tmp = path.with_name(path.name + ".tmp")
    try:  # write-then-replace: a crash mid-write never leaves a truncated cache
        tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(tmp, path)
    except OSError as exc:
        tmp.unlink(missing_ok=True)
        print(f"[pod] warning: cannot write cache: {exc}", file=sys.stderr)


_NOUN = {"refrigerator": "refrigerator", "washer": "laundry (washer/dryer)", "cooking": "cooking appliance"}


def _prompt(texts: list[str], category: str = DEFAULT_MAJOR) -> str:
    keys = "\n".join(f"- {k}: {en}" for k, _, en, _ in TAXONOMIES[category])
    items = "\n".join(f"{i}: {t}" for i, t in enumerate(texts))
    return (f"Classify each {_NOUN[category]} feature into exactly ONE key from this list. "
            f"Use 'other' if none fits.\n{keys}\n\nFeatures:\n{items}\n\n"
            'Reply with JSON only: an object mapping the feature number to the key, e.g. {"0": "led_light"}.')


def _llm_classify(texts: list[str], llm_fn: Callable, category: str = DEFAULT_MAJOR) -> dict[int, str]:
    """One batched call -> {index: valid_key}; missing/invalid indices are absent."""
    tax = _TAXES[category]
    try:
        reply = llm_fn(_prompt(texts, category))
    except Exception as exc:  # noqa: BLE001 - LLM failure must never break the run
        print(f"[pod] warning: LLM call failed ({exc}); using 'other'", file=sys.stderr)
        return {}
    if reply is None:
        print("[pod] warning: LLM unavailable; unmatched items -> 'other'", file=sys.stderr)
        return {}
    pairs: list[tuple[object, object]] = []
    if isinstance(reply, dict):
        pairs = list(reply.items())
    elif isinstance(reply, list):
        for e in reply:
            if isinstance(e, dict):
                pairs.append((e.get("index", e.get("id", e.get("i"))), e.get("key", e.get("taxonomy_key"))))
    out: dict[int, str] = {}
    for i, k in pairs:
        try:
            idx = int(i)
        except (TypeError, ValueError):
            continue
        if 0 <= idx < len(texts) and isinstance(k, str) and k in tax:
            out[idx] = k
    return out


def _item(product: ProductRecord, raw: str, key: str, by: str, major: str = DEFAULT_MAJOR) -> PodItem:
    _, cat, en, ko = _TAXES[major][key]
    return PodItem(brand=product.brand, model_number=product.model_number, raw_text=raw,
                   taxonomy_key=key, category=cat, name_en=en, name_ko=ko, matched_by=by)


def _resolve(units: list[tuple[str, str]], llm_fn: Optional[Callable],
             cache_path: Path, category: str = DEFAULT_MAJOR) -> list[list[tuple[str, str]]]:
    """units = (keyword_text, llm_text). Returns per unit a list of (key, matched_by)."""
    out: list[list[tuple[str, str]]] = [[] for _ in units]
    pending: list[int] = []
    for i, (kw, _) in enumerate(units):
        ks = keyword_keys(kw, category)
        if ks:
            out[i] = [(k, "keyword") for k in ks]
        else:
            pending.append(i)
    if not pending:
        return out
    cache = _load_cache(cache_path)
    tax = _TAXES[category]
    todo: list[int] = []
    for i in pending:
        cached = cache.get(_sha(units[i][1], category))
        if cached in tax:
            out[i] = [(cached, "llm")]
        else:
            todo.append(i)
    if todo:
        if llm_fn is None:
            import llm
            llm_fn = llm.chat_json
        answers = _llm_classify([units[i][1] for i in todo], llm_fn, category)
        for n, i in enumerate(todo):
            if n in answers:
                out[i] = [(answers[n], "llm")]
                cache[_sha(units[i][1], category)] = answers[n]
            else:
                out[i] = [(OTHER, "none")]
        if answers:
            _save_cache(cache_path, cache)
    return out


def _norm_id(s: Optional[str]) -> str:
    return (s or "").strip().casefold()


def _doc_label(source_doc: str) -> str:
    name = Path(source_doc or "").stem
    return "Manual" if "manual" in name.lower() else (name or "doc")


def _page(m) -> str:
    return f" p{m.source_page}" if m.source_page else ""


def normalize_pod(product: ProductRecord, llm_fn: Optional[Callable] = None,
                  cache_path: Optional[Path] = None, modes: Optional[list] = None) -> list[PodItem]:
    """Map pod features (and optional ModeRecords) to taxonomy keys.

    A raw feature may yield several items (one per matched key, same raw_text). Mode-derived
    items (matched_by='mode') are only added for keys not already evidenced by pod_features;
    otherwise the mode wording is attached as mode_ref on the existing item.
    """
    cache_path = Path(cache_path) if cache_path else CACHE_PATH
    major = major_of_product(product)  # taxonomy follows the product's major group
    raws: list[str] = []
    seen: set[str] = set()
    for r in product.pod_features:
        r = (r or "").strip()
        if r and r.lower() not in seen:
            seen.add(r.lower())
            raws.append(r)

    my_modes = []
    seen_modes: set[str] = set()
    for m in modes or []:
        name = _norm_id(m.mode_name)
        if (_norm_id(m.brand) == _norm_id(product.brand)
                and _norm_id(m.model_number) == _norm_id(product.model_number)
                and name and name not in seen_modes):
            seen_modes.add(name)
            my_modes.append(m)

    units = [(r, r) for r in raws]
    units += [(m.mode_name, f"{m.mode_name}: {(m.description or '')[:150]}") for m in my_modes]
    resolved = _resolve(units, llm_fn, cache_path, major) if units else []

    items: list[PodItem] = []
    for r, res in zip(raws, resolved):
        for key, by in res:
            items.append(_item(product, r, key, by, major))

    for m, res in zip(my_modes, resolved[len(raws):]):
        raw = f"{m.mode_name} (Modes: {_doc_label(m.source_doc)}{_page(m)})"
        ref = f"Modes: {m.mode_name}{_page(m)}"
        for key, _ in res:
            existing = next((i for i in items if i.taxonomy_key == key), None) if key != OTHER else None
            if existing is None:
                items.append(_item(product, raw, key, "mode", major))
            else:
                existing.mode_ref = " | ".join(x for x in (existing.mode_ref, ref) if x)

    present = {i.taxonomy_key for i in items}
    for attr, key, text in _FIELD_EVIDENCE_BY[major]:
        if getattr(product, attr, None) is True and key not in present:
            items.append(_item(product, text, key, "keyword", major))
            present.add(key)
    return items


# ---------------------------------------------------------------- Excel
_HDR_FILL = PatternFill("solid", fgColor="1F4E78")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_GREEN = PatternFill("solid", fgColor="C6EFCE")
_GREY = PatternFill("solid", fgColor="D9D9D9")
_AMBER = PatternFill("solid", fgColor="FFE699")


def _put(ws, row: int, col: int, value):
    value = clean_text(value)
    c = ws.cell(row=row, column=col)
    c.value = value
    if is_formula_like(value):
        c.data_type = "s"
        c.quotePrefix = True  # Excel keeps it as text even if the user edits the cell
    return c


def _style_header(ws, ncols: int):
    for c in ws[1][:ncols]:
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def _fit(ws, cap: int = 70):
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), cap)


def _fresh_sheet(wb: Workbook, title: str):
    if title in wb.sheetnames:
        del wb[title]
    return wb.create_sheet(title)


def _wording(i: PodItem) -> str:
    return " | ".join(x for x in (i.raw_text, i.mode_ref) if x)


def _source(its: list[PodItem]) -> str:
    web = any(i.matched_by != "mode" and not i.raw_text.endswith("(spec field)") for i in its)
    mode = any(i.matched_by == "mode" or i.mode_ref for i in its)
    field = any(i.raw_text.endswith("(spec field)") for i in its)
    if web and mode:
        return "both"
    return "web" if web else "modes" if mode else "spec field" if field else ""


def compare_rows(products: list[ProductRecord], items_by_product: list[list[PodItem]],
                 category: Optional[str] = None) -> list[dict]:
    """Matrix rows in the major group's taxonomy order; rows where no product has the item are omitted.
    Refuses to mix major categories (compare washers with washers, never with fridges)."""
    majors = {major_of_product(p) for p in products}
    if len(majors) > 1:
        raise ValueError(f"cannot compare across major categories: {sorted(majors)}")
    major = category or (majors.pop() if majors else DEFAULT_MAJOR)
    brands = [p.brand for p in products]
    rows = []
    for key, cat, en, ko in TAXONOMIES[major]:
        matched = [[i for i in its if i.taxonomy_key == key] for its in items_by_product]
        flags = [bool(m) for m in matched]
        if not any(flags):
            continue
        diff = "Both" if all(flags) else "Only " + ", ".join(b for b, f in zip(brands, flags) if f)
        rows.append({"key": key, "category": cat, "en": en, "ko": ko, "present": flags,
                     "wording": [" | ".join(dict.fromkeys(_wording(i) for i in m)) for m in matched],
                     "source": [_source(m) for m in matched], "diff": diff})
    return rows


def add_pod_sheets(wb: Workbook, products: list[ProductRecord],
                   items_by_product: Optional[list[list[PodItem]]] = None,
                   llm_fn: Optional[Callable] = None, modes: Optional[list] = None) -> None:
    """Add 'POD_Items' and the POD comparison sheet(s) to wb (not saved).

    One comparison matrix per major category: 'POD_Compare' when a single major group is present, else
    'POD_Compare_<category>' per group (major categories are never mixed in one matrix)."""
    if items_by_product is None:
        items_by_product = [normalize_pod(p, llm_fn=llm_fn, modes=modes) for p in products]

    ws = _fresh_sheet(wb, "POD_Items")
    headers = ["Brand", "Model", "Raw text", "Taxonomy key", "Category", "Name (EN)", "Name (KO)", "Matched by", "Mode ref"]
    ws.append(headers)
    for its in items_by_product:
        for i in its:
            r = ws.max_row + 1
            vals = [i.brand, i.model_number, i.raw_text, i.taxonomy_key, i.category,
                    i.name_en, i.name_ko, i.matched_by, i.mode_ref]
            for c, v in enumerate(vals, 1):
                _put(ws, r, c, v)
    _style_header(ws, len(headers))
    _fit(ws)

    by_major: dict[str, list[int]] = {}
    for i, p in enumerate(products):
        by_major.setdefault(major_of_product(p), []).append(i)
    ordered = [m for m in TAXONOMIES if m in by_major] or [DEFAULT_MAJOR]
    for stale in [n for n in wb.sheetnames if n == "POD_Compare" or n.startswith("POD_Compare_")]:
        del wb[stale]
    for major in ordered:  # one matrix per major group; plain 'POD_Compare' when only one group is present
        idx = by_major.get(major, [])
        _compare_sheet(wb, "POD_Compare" if len(ordered) == 1 else f"POD_Compare_{major}",
                       [products[i] for i in idx], [items_by_product[i] for i in idx], major)


def _compare_sheet(wb: Workbook, title: str, products: list[ProductRecord],
                   items_by_product: list[list[PodItem]], major: str) -> None:
    wc = wb.create_sheet(title)
    head = ["Korean name", "English name", "Category"]
    for p in products:
        head += [f"{p.brand} present", f"{p.brand} wording", f"{p.brand} source"]
    head.append("Difference")
    wc.append(head)
    for row in compare_rows(products, items_by_product, major):
        r = wc.max_row + 1
        vals = [row["ko"], row["en"], row["category"]]
        for f, w, src in zip(row["present"], row["wording"], row["source"]):
            vals += ["✓" if f else "–", w, src]
        vals.append(row["diff"])
        for c, v in enumerate(vals, 1):
            _put(wc, r, c, v)
        for n, f in enumerate(row["present"]):
            cell = wc.cell(row=r, column=4 + 3 * n)
            cell.fill = _GREEN if f else _GREY
            cell.alignment = Alignment(horizontal="center")
        if row["diff"].startswith("Only"):
            for c in (1, 2, 3, len(vals)):
                wc.cell(row=r, column=c).fill = _AMBER
    _style_header(wc, len(head))
    _fit(wc, cap=60)
