"""Excel export of a competitor-match result (match.rank + match.launches), built in memory.

Sheets: 요약 (target, weights, counts, data readiness), 순위 (one row per ranked model with every score component and
its evidence), 가격 5단계, 신제품 (per brand, with the evidence), 신제품 트렌드. All strings go through excel_writer's
sanitising helpers (control characters dropped, formula-looking text forced to text).
"""
import io
from datetime import datetime

from openpyxl import Workbook

import catalog
from excel_writer import _add_table

BASIS_KO = {"release_date": "출시일", "site_new": "사이트 NEW", "first_seen": "최초 발견", "site_order": "사이트 최신순",
            "unknown": "출시 시점 미확인"}
FEATURE_KO = {"wifi": "Wi-Fi", "convection": "컨벡션", "air_fry": "에어프라이", "steam": "스팀", "energy_star": "ENERGY STAR"}
SPEC_KO = {"width_in": "폭(in)", "capacity_total_cuft": "용량(cu ft)", "oven_capacity_cuft": "오븐 용량(cu ft)", "fuel": "열원",
           "burners": "버너 수"}


def _pct(score):
    """0..1 score -> 0..100 (None stays empty = unknown)."""
    return None if score is None else round(score * 100)


def _spec_text(details: list[dict]) -> str:
    return "; ".join(f"{SPEC_KO.get(d['key'], FEATURE_KO.get(d['key'], d['key']))}: 기준 {d['want']} / 모델 "
                     f"{'미확인' if d['actual'] is None else d['actual']}" for d in details)


def _pad(rows: list[list], width: int) -> list[list]:
    """Every row as long as the header (the sheet helper reads one cell per column from every row)."""
    return [r + [None] * (width - len(r)) for r in rows]


def _summary_rows(out: dict, readiness: dict, window: int) -> list[list]:
    t, c = out["target"], out["counts"]
    specs = t.get("specs") or {}
    spec_txt = "; ".join(f"{SPEC_KO.get(k, k)}={'/'.join(map(str, v)) if isinstance(v, list) else v}"
                         for k, v in specs.items() if v not in (None, [], "")) or "입력 없음"
    rows = [
        ["생성 시각", f"{datetime.now():%Y-%m-%d %H:%M}"],
        ["소분류", catalog.label_ko(t["sub"])],
        ["국가", t["country"]],
        ["통화", t["currency"]],
        ["목표 가격", t.get("price")],
        ["기준 가격 단계", f"{t['tier']}단계 ({t['tier_label']})" if t.get("tier") else "미확인"],
        ["가격 허용폭(±%)", out["band_pct"]],
        ["입력 스펙", spec_txt],
        ["중요도 가중치", "; ".join(f"{k}={v:g}" for k, v in out["weights"].items())],
        ["단계 산정 범위", out["tier_scope"]],
        ["후보 수", c["pool"]], ["순위에 표시", c["ranked"]], ["폭 조건 밖 제외", c["excluded_by_spec"]],
        ["단계 범위 밖 제외", c["outside_tier_window"]], ["가격 미확인(순위 제외)", c["price_unknown"]],
        ["분석 대상 모델", readiness["models"]], ["가격 확인", readiness["with_price"]], ["평점 확인", readiness["with_rating"]],
        ["출시일 확인", readiness["with_release_date"]], ["사이트 NEW 표시", readiness["site_new_flagged"]],
        ["앱이 새로 발견한 모델", readiness["new_discoveries"]],
        ["신제품 판단 기간(개월)", window],
    ]
    if out.get("distrusted_new_flags"):
        rows.append(["신뢰하지 않은 NEW 표시", ", ".join(out["distrusted_new_flags"])])
    rows.append(["주의", "점수는 참고용입니다. 최근성은 사이트 표시·날짜·앱의 최초 발견 시점에서 추정한 값이며, 데이터가 없는 항목은 "
                       "중립 점수로 계산하고 '근거 비율'로 실제 데이터 비중을 보여 줍니다."])
    return rows


RANK_HEADER = ["순위", "브랜드", "모델", "제품명", "가격", "통화", "단계", "단계명", "단계 차이(기준 대비)", "종합 점수", "근거 비율(%)",
               "가격 점수", "목표 대비(%)", "스펙 점수", "스펙 상세", "최근성 점수", "최근성 근거", "근거 설명", "출시 후 개월",
               "신제품", "평점", "리뷰 수", "보정 평점", "호응 점수", "URL"]


def _ranking_rows(out: dict) -> list[list]:
    rows = []
    for i, r in enumerate(out["results"], 1):
        comp = r["components"]
        rec, resp = comp["recency"], comp["response"]
        rows.append([
            i, r["brand"], r["model_number"], r["name"], r["price"], r["currency"],
            r["tier"], r["tier_label"], r["tier_diff"],
            None if r["total"] is None else round(r["total"]), round(r["coverage"] * 100),
            _pct(comp["price"]["score"]), comp["price"]["delta_pct"],
            _pct(comp["spec"]["score"]), _spec_text(comp["spec"]["details"]),
            _pct(rec["score"]), BASIS_KO.get(rec["basis"], rec["basis"]), rec["evidence"], rec["months"],
            "예" if r["is_launch"] else "아니오",
            resp.get("rating"), resp.get("reviews"), resp.get("adjusted"), _pct(resp["score"]),
            r["url"]])
    return rows


def build(out: dict, launches: dict, readiness: dict) -> bytes:
    """xlsx bytes of one competitor-match result. `readiness` is server._data_readiness of the same pool."""
    wb = Workbook()
    wb.remove(wb.active)
    _add_table(wb, "요약", ["항목", "값"], _summary_rows(out, readiness, launches["window_months"]))
    _add_table(wb, "순위", RANK_HEADER, _ranking_rows(out))
    _add_table(wb, "가격 5단계", ["단계", "이름", "최저 가격", "최고 가격", "모델 수"],
               [[t["tier"], t["label"], t["min"], t["max"], t["count"]] for t in out["tiers"]])
    cur = out["target"]["currency"]
    new_head = ["브랜드", "모델", "제품명", "가격", "통화", "근거", "근거 설명", "평점", "리뷰 수", "URL"]
    new_rows = [[brand, m["model_number"], m["name"], m["price"], m.get("currency") or cur, BASIS_KO.get(m["basis"], m["basis"]),
                 m["evidence"], m.get("rating"), m.get("review_count"), m["url"]]
                for brand, models in launches["by_brand"].items() for m in models]
    _add_table(wb, "신제품", new_head, new_rows)
    med = launches["median_price"]
    trend_head = ["기능", "신제품 비율(%)", "신제품 표본", "기존 비율(%)", "기존 표본", "차이(%p)", "표본 적음"]
    trend = [[FEATURE_KO.get(t["feature"], t["feature"]), t["new"]["pct"], t["new"]["known"], t["existing"]["pct"],
              t["existing"]["known"], t["delta_pts"], "예" if t["low_sample"] else "아니오"] for t in launches["trend"]]
    trend += [[None], ["가격 중앙값 (신제품)", med["new"]], ["가격 중앙값 (기존)", med["existing"]], ["참고", launches["note"]]]
    _add_table(wb, "신제품 트렌드", trend_head, _pad(trend, len(trend_head)))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
