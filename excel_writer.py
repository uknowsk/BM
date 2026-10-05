"""Write scraped records to a multi-sheet Excel workbook."""
import os
import sys
from datetime import datetime
from pathlib import Path
from types import UnionType
from typing import Callable, Iterable, Optional, Union, get_args, get_origin

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from compare_model import HEADER_IDS, build_compare, group_products
from schema import DocumentRecord, ModeRecord, ProductRecord, RawSpec

ROOT = Path(__file__).parent
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
LINK_FONT = Font(color="0563C1", underline="single")

# Korean descriptions and units, keyed by field name (shared names mean the same thing).
KO = {
    "brand": "브랜드", "model_number": "모델 번호", "product_name": "제품명", "category": "제품 대분류 (refrigerator/washer/cooking 키 또는 기존 라벨)",
    "subcategory": "제품 소분류 키 (예: french_door, front_load, induction)",
    "extra_specs": "제품군별 추가 사양 (영문 항목명: 값, '; '로 구분 — 상세는 Category_Specs 시트)",
    "door_style": "도어 스타일 (냉장고 전용)", "finish_color": "마감 색상", "product_url": "제품 페이지 URL",
    "price_usd": "가격 (미국 달러)", "capacity_total_cuft": "총 용량 (냉장고 전용; 다른 제품군은 extra_specs)", "capacity_fridge_cuft": "냉장실 용량 (냉장고 전용)",
    "capacity_freezer_cuft": "냉동실 용량 (냉장고 전용)", "width_in": "폭", "height_in": "높이", "depth_in": "깊이",
    "weight_lb": "무게", "voltage_v": "전압", "amps": "전류", "frequency_hz": "주파수",
    "energy_kwh_year": "연간 에너지 소비량", "energy_star": "에너지 스타 인증 여부",
    "ice_maker": "제빙기 탑재 여부 (냉장고 전용)", "water_dispenser": "정수(물) 디스펜서 탑재 여부 (냉장고 전용)",
    "wifi_supported": "Wi-Fi 지원 여부 (추론)", "wifi_evidence": "Wi-Fi 추론 근거 문장/키워드",
    "pod_features": "주요 특징 (영문, '; '로 구분)", "scraped_at": "수집 일시 (ISO 8601)",
    "doc_type": "문서 유형 (Manual, QuickSpecs, EnergyGuide, Installation, Warranty, SpecSheet, Other)",
    "source_url": "문서 원본 URL", "local_path": "로컬 저장 경로", "sha256": "파일 SHA-256 해시",
    "size_bytes": "파일 크기", "pages": "페이지 수", "downloaded_at": "다운로드 일시 (ISO 8601)",
    "source": "출처 (web 또는 PDF 파일명)", "section": "원문 섹션명", "key": "사양 항목명 (원문)",
    "value": "사양 값 (원문)",
    "image_url": "제품 대표 이미지 원본 URL (https)", "image_path": "내려받은 대표 이미지 경로 (프로젝트 루트 기준)",
    "region": "판매 지역 (catalog.REGIONS 키)", "country": "국가 코드", "currency": "통화 코드 (ISO 4217)",
    "price_local": "현지 통화 가격",
    "fridge_temp_range_f": "냉장실 설정 온도 범위 (매뉴얼에 명시된 경우만)",
    "freezer_temp_range_f": "냉동실 설정 온도 범위 (매뉴얼에 명시된 경우만)",
    "mode_name": "모드/기능 이름 (원문)",
    "description": "모드/기능 설명 (매뉴얼 근거)",
    "setting_range": "설정 범위 (예: 34-44 F)", "how_to_activate": "활성화(설정) 방법",
    "source_doc": "출처 문서 파일명", "source_page": "출처 페이지 번호 (1부터)",
}
KO_BY_SHEET = {("Modes", "category"): "모드 분류 (Cooling, Freezing, Freshness, Away, Religious, Energy, Ice&Water, Alert, Smart, Other)"}
UNITS = {
    "price_usd": "USD", "capacity_total_cuft": "cu ft", "capacity_fridge_cuft": "cu ft",
    "capacity_freezer_cuft": "cu ft", "width_in": "in", "height_in": "in", "depth_in": "in",
    "weight_lb": "lb", "voltage_v": "V", "amps": "A", "frequency_hz": "Hz",
    "energy_kwh_year": "kWh/yr", "size_bytes": "bytes",
}
SHEETS = (("Products", ProductRecord), ("Documents", DocumentRecord), ("Raw_Specs", RawSpec), ("Modes", ModeRecord))


IMAGE_COL_NAME = "image"  # trailing Products column holding the embedded picture (not a ProductRecord field)
IMAGE_PX = 120  # picture width in the sheet
HEADER_IDS_SET = set(HEADER_IDS)

CATEGORY_SPECS_HEADER = ["brand", "model_number", "product_name", "category", "subcategory", "key", "value"]
CATEGORY_SPECS_KO = {
    "brand": "브랜드", "model_number": "모델 번호", "product_name": "제품명", "category": "제품 대분류",
    "subcategory": "제품 소분류", "key": "제품군별 사양 항목명 (영문)", "value": "사양 값",
}


def _type_name(annotation) -> str:
    if get_origin(annotation) in (Union, UnionType):
        args = [a for a in get_args(annotation) if a is not type(None)]
        return " | ".join(_type_name(a) for a in args) + " (optional)"
    if get_origin(annotation) is dict:
        k, v = get_args(annotation)
        return f"dict[{_type_name(k)}, {_type_name(v)}]"
    if get_origin(annotation) is list:
        return f"list[{_type_name(get_args(annotation)[0])}]"
    return getattr(annotation, "__name__", str(annotation))


def _cell_value(value):
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    return "; ".join(map(str, value)) if isinstance(value, list) else value


MAX_CELL_CHARS = 32767  # Excel's hard per-cell limit


def clean_text(value):
    """Strings only: drop control characters openpyxl refuses to write and cap at Excel's cell limit."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_CHARS]
    return value


_FORMULA_CHARS = ("=", "+", "-", "@", "\t", "\r")


def is_formula_like(value) -> bool:
    """True for strings Excel/CSV consumers could evaluate: first (or first non-blank) char in = + - @ TAB CR."""
    return isinstance(value, str) and (value[:1] in _FORMULA_CHARS or value.lstrip()[:1] in _FORMULA_CHARS)


def _neutralize(cell) -> None:
    if is_formula_like(cell.value):
        cell.data_type = "s"
        cell.quotePrefix = True  # Excel keeps it as text even if the user edits the cell


def _put(ws, row: list) -> None:
    """Append a row; scraped strings that look like formulas are forced to text (no formula injection)."""
    ws.append([clean_text(v) for v in row])
    for cell in ws[ws.max_row]:
        _neutralize(cell)


def _style_header(ws, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _fit_widths(header: list[str], rows: list[list], cap: int = 60) -> list[int]:
    widths = []
    for i, h in enumerate(header):
        longest = max([len(str(h))] + [len(str(r[i])) for r in rows if r[i] is not None])
        widths.append(min(max(longest + 2, 10), cap))
    return widths


def _add_table(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        _put(ws, r)
    _style_header(ws, _fit_widths(header, rows))


def _model_rows(model: type[BaseModel], records: Iterable[BaseModel]) -> list[list]:
    fields = list(model.model_fields)
    return [[_cell_value(getattr(rec, f)) for f in fields] for rec in records]


def write_excel(products: Iterable[ProductRecord], documents: Iterable[DocumentRecord],
                raw_specs: Iterable[RawSpec], path: str | Path,
                run_log: Iterable[tuple[str, str, str]] | None = None,
                modes: Iterable[ModeRecord] | None = None,
                extra_sheets: Callable[[Workbook], None] | None = None,
                pod_items: list | None = None) -> Path:
    """Write the workbook to `path` (parent dirs created). Optional run_log adds a 'Run_Log' sheet;
    extra_sheets(wb) may add further sheets (e.g. POD comparison) before saving. The transposed
    'Compare' sheet(s) come first; pod_items (per-product PodItem lists) adds its POD rows."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = {"Products": products, "Documents": documents, "Raw_Specs": raw_specs, "Modes": modes or []}

    products, raw_specs, documents = list(products), list(raw_specs), list(documents)
    modes = list(modes or [])
    data.update(Products=products, Documents=documents, Raw_Specs=raw_specs, Modes=modes)
    wb = Workbook()
    wb.remove(wb.active)
    for title, model in SHEETS:
        header, rows = list(model.model_fields), _model_rows(model, data[title])
        if title == "Products":
            header, rows = header + [IMAGE_COL_NAME], [r + [None] for r in rows]
        _add_table(wb, title, header, rows)
    try:
        _embed_images(wb["Products"], products, len(ProductRecord.model_fields) + 1, ROOT)
    except Exception as exc:  # noqa: BLE001 - Pillow missing / odd image: the workbook is still useful without pictures
        print(f"[excel] images not embedded: {type(exc).__name__}", file=sys.stderr)
    _add_table(wb, "Category_Specs", CATEGORY_SPECS_HEADER,
               [[p.brand, p.model_number, p.product_name, p.category, p.subcategory, k, v]
                for p in products for k, v in p.extra_specs.items()])

    if products:
        _add_compare_sheets(wb, products, documents, raw_specs, modes, pod_items, ROOT)

    docs_ws = wb["Documents"]
    downloads = (ROOT / "downloads").resolve()
    col = list(DocumentRecord.model_fields).index("local_path") + 1
    for row in range(2, docs_ws.max_row + 1):
        cell = docs_ws.cell(row, col)
        if cell.value:
            target = Path(cell.value)
            resolved = (target if target.is_absolute() else ROOT / target).resolve()
            if not resolved.is_relative_to(downloads):  # only link files we downloaded, never arbitrary paths
                continue
            try:  # relative to the workbook: no absolute file:/// URL (which embeds the OS username)
                rel = os.path.relpath(resolved, path.resolve().parent)
            except ValueError:  # different drive: no portable relative link
                continue
            cell.hyperlink = rel.replace(os.sep, "/")
            cell.font = LINK_FONT

    schema_rows = [
        [name, _type_name(info.annotation), info.description or "", UNITS.get(name, ""), KO_BY_SHEET.get((sheet, name)) or KO.get(name, ""), sheet]
        for sheet, model in SHEETS for name, info in model.model_fields.items()
    ]
    schema_rows += [[name, "str", "", "", CATEGORY_SPECS_KO[name], "Category_Specs"] for name in CATEGORY_SPECS_HEADER]
    _add_table(wb, "Schema", ["field", "type", "description", "unit", "description_ko", "sheet"], schema_rows)

    if extra_sheets is not None:
        extra_sheets(wb)

    if run_log is not None:
        _add_table(wb, "Run_Log", ["brand", "status", "message"], [list(r) for r in run_log])

    return _atomic_save(wb, path)


def _embed_images(ws, products: list[ProductRecord], col: int, root: Path) -> None:
    """Anchor each product's downloaded picture in its Products row (column `col`, ~IMAGE_PX wide) and grow the row
    to fit. Only files inside downloads/ are read; an unreadable file just leaves that cell empty."""
    from openpyxl.drawing.image import Image as XlImage  # needs Pillow: ImportError is handled by the caller
    downloads = (root / "downloads").resolve()
    letter = get_column_letter(col)
    placed = False
    for i, p in enumerate(products):
        if not p.image_path:
            continue
        path = (root / p.image_path).resolve()
        if not (path.is_relative_to(downloads) and path.is_file()):
            continue
        try:
            img = XlImage(str(path))
            scale = IMAGE_PX / img.width
            img.width, img.height = IMAGE_PX, max(1, round(img.height * scale))
        except Exception:  # noqa: BLE001 - corrupt file
            continue
        row = i + 2
        ws.add_image(img, f"{letter}{row}")
        ws.row_dimensions[row].height = img.height * 0.75 + 4  # px -> points
        placed = True
    if placed:
        ws.column_dimensions[letter].width = IMAGE_PX / 7 + 2


SECTION_FILL = {"기본정보": "D9E2F3", "치수·무게": "E2EFDA", "용량": "DDEBF7", "전기·에너지": "FFF2CC", "성능": "FCE4D6", "기능": "E4DFEC",
                "디자인": "EDEDED", "연결성": "D0E8F0", "조리(오븐·쿡탑)": "F8E1D0", "세탁·건조": "D9EAD3", "냉각·신선": "CFE2F3", "보증·기타": "E7E6E6"}
UNSURE_FILL = PatternFill("solid", fgColor="FFE699")  # label cell of a row whose canonical mapping needs review
QUIET_FONT = Font(color="8C8C8C")                      # rows that are identical for every product
LONG_TAIL = "롱테일 항목 (펼쳐서 보기)"
YES_FILL, NO_FILL, DIFF_MARK = (PatternFill("solid", fgColor=c) for c in ("C6EFCE", "EDEDED", "FFC000"))
HEAD_PX = 140          # header picture box
HEADER_ROWS = 6        # rows 1-5 = image / brand / model / name / price, row 6 = column titles
FIRST_PRODUCT_COL = 5  # A 구분, B 항목(KO), C 항목(EN), D 단위, then one column per product
NL = chr(10)


def _picture(ws, rel: Optional[str], anchor: str, root: Path) -> int:
    """Embed the picture (fit inside HEAD_PX x HEAD_PX) at `anchor`; returns its pixel height (0 if none)."""
    if not rel:
        return 0
    from openpyxl.drawing.image import Image as XlImage
    downloads = (root / "downloads").resolve()
    path = (root / rel).resolve()
    if not (path.is_relative_to(downloads) and path.is_file()):
        return 0
    try:
        img = XlImage(str(path))
        k = min(HEAD_PX / img.width, HEAD_PX / img.height, 1.0)
        img.width, img.height = max(1, round(img.width * k)), max(1, round(img.height * k))
    except Exception:  # noqa: BLE001 - corrupt file
        return 0
    ws.add_image(img, anchor)
    return img.height


def _cell_for(row: dict, value, note) -> tuple:
    """(cell value, fill, font) for one product value of a compare row."""
    if row["kind"] == "flag":
        if value:
            return ("✓" + (NL + str(note) if note else ""), YES_FILL, Font(color="1F6B47", bold=True))
        return ("–", NO_FILL, Font(color="808080"))
    if value is None or value == "":
        return ("–", NO_FILL, Font(color="808080"))
    return (value, None, None)


def _source_comment(row: dict, sources: list) -> str:
    """Cell note: the product's ORIGINAL label / value (and how it was mapped) when that is not just the canonical wording."""
    if not sources:
        return ""
    plain = (len(sources) == 1 and sources[0]["method"] in ("seed", "override")
             and sources[0]["label"].strip().lower() in (str(row["key_en"]).lower(), str(row["key_ko"]).lower())
             and str(sources[0]["value"]).strip().lower() in ("yes", "no", ""))
    if plain:
        return ""
    lines = ["원문 항목명 → 값 (매핑)"]
    for s in sources:
        via = f" [{s['via']}]" if s.get("via") else ""
        lines.append(f"{s['label']}{via} = {s['value']}   ({s['method']} {s['score']:.2f})")
    return NL.join(lines)


def _add_compare_sheet(wb: Workbook, index: int, title: str, ps: list[ProductRecord], rows: list[dict], root: Path) -> None:
    """Transposed comparison: items are rows, products are columns (see compare_model)."""
    from openpyxl.worksheet.properties import Outline, PageSetupProperties
    ws = wb.create_sheet(title, index)
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False)  # the band above a block is its summary row
    n = len(ps)
    last, diffc = FIRST_PRODUCT_COL + n - 1, FIRST_PRODUCT_COL + n
    by_id = {r["id"]: r for r in rows}
    labels = {1: "이미지", 2: "브랜드", 3: "모델", 4: "제품명", 5: "가격 (USD)"}
    for r, text in labels.items():
        c = ws.cell(r, 2, text)
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center")
    for i, p in enumerate(ps):
        col = FIRST_PRODUCT_COL + i
        for r, v in ((2, p.brand), (3, p.model_number), (4, p.product_name), (5, p.price_usd)):
            c = ws.cell(r, col, clean_text(v))
            c.font = Font(bold=r in (2, 3), size=12 if r == 3 else 11)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            _neutralize(c)
        ws.cell(5, col).number_format = '"$"#,##0'
    heights = [_picture(ws, p.image_path, f"{get_column_letter(FIRST_PRODUCT_COL + i)}1", root) for i, p in enumerate(ps)]
    ws.row_dimensions[1].height = (max(heights) + 8) * 0.75 if any(heights) else 24
    for c in range(1, diffc + 1):
        t = ws.cell(HEADER_ROWS, c, ["구분", "항목", "Item (EN)", "단위"][c - 1] if c < FIRST_PRODUCT_COL
                    else "diff" if c == diffc else clean_text(f"{ps[c - FIRST_PRODUCT_COL].brand} {ps[c - FIRST_PRODUCT_COL].model_number}"))
        t.font, t.fill, t.alignment = HEADER_FONT, HEADER_FILL, Alignment(vertical="center", wrap_text=True)
    section, tail_band = None, False
    for row in rows:
        if row["id"] in HEADER_IDS_SET:
            continue
        if row["section"] != section:
            section, tail_band = row["section"], False
            r = ws.max_row + 1
            for c in range(1, last + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=SECTION_FILL.get(section, "DDDDDD"))
            ws.cell(r, 1, section).font = Font(bold=True, size=12)
            ws.cell(r, diffc, "band")
        long_tail = not row.get("core", True)
        if long_tail and not tail_band:  # long-tail rows of a section live in one collapsed outline group
            tail_band = True
            r = ws.max_row + 1
            ws.cell(r, 1, section)
            g = ws.cell(r, 2, LONG_TAIL)
            g.font = Font(bold=True, italic=True, color="404040")
            for c in range(1, last + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F2F2F2")
            ws.row_dimensions[r].outline_level = 1
            ws.row_dimensions[r].collapsed = True
            ws.cell(r, diffc, "band")
        r = ws.max_row + 1
        ws.cell(r, 1, section)
        for c, v in ((2, row["key_ko"]), (3, row["key_en"]), (4, row["unit"])):
            cell = ws.cell(r, c, clean_text(v))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            _neutralize(cell)
            if row.get("same"):
                cell.font = QUIET_FONT
        if row.get("uncertain"):
            ws.cell(r, 3).fill = UNSURE_FILL
        srcs = row.get("sources") or [[] for _ in row["values"]]
        for i, v in enumerate(row["values"]):
            val, fill, font = _cell_for(row, v, row["notes"][i])
            cell = ws.cell(r, FIRST_PRODUCT_COL + i, clean_text(val))
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if row["kind"] == "flag" or fill else "left")
            if fill:
                cell.fill, cell.font = fill, font
            elif row.get("same"):
                cell.font = QUIET_FONT
            _neutralize(cell)
            text = _source_comment(row, srcs[i] if i < len(srcs) else [])
            if text:
                com = Comment(clean_text(text), "canon")
                com.width, com.height = 360, min(60 + 18 * text.count(NL), 340)
                cell.comment = com
        if row["differs"]:
            ws.cell(r, 1).fill = DIFF_MARK
        ws.cell(r, diffc, "diff" if row["differs"] else "")
        ws.row_dimensions[r].outline_level = 2 if long_tail else 1
        ws.row_dimensions[r].hidden = long_tail
    widths = [10, 26, 40, 9] + [30] * n + [6]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions[get_column_letter(diffc)].hidden = True
    ws.freeze_panes = ws.cell(HEADER_ROWS + 1, FIRST_PRODUCT_COL)
    ws.auto_filter.ref = f"A{HEADER_ROWS}:{get_column_letter(diffc)}{max(ws.max_row, HEADER_ROWS + 1)}"
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"1:{HEADER_ROWS}"


def _add_compare_sheets(wb: Workbook, products, documents, raw_specs, modes, pod_items, root: Path) -> None:
    """First sheet(s): 'Compare' (one major category) or 'Compare_<major>' per category, then the 'Mapping' audit sheet."""
    model = build_compare(products, documents, raw_specs, modes, pod_items)
    groups = group_products(products)
    for i, (major, rows) in enumerate(model.items()):
        _add_compare_sheet(wb, i, "Compare" if len(model) == 1 else f"Compare_{major}"[:31], groups[major], rows, root)
    _add_mapping_sheet(wb, len(model), model, groups)


MAPPING_PER_PRODUCT = ("원문 항목명", "원문 값", "방법", "점수")


def _add_mapping_sheet(wb: Workbook, index: int, model: dict, groups: dict) -> None:
    """'Mapping' (audit): one row per canonical attribute with every product's source label / value / method / score and a
    needs-review flag (method llm, or an embedding merge scored < canon.REVIEW_BELOW). Fix a wrong mapping with data/canon_overrides.json."""
    import canon
    ws = wb.create_sheet("Mapping", index)
    fixed = ["구분(major)", "canonical id", "항목 (EN)", "항목 (KO)", "섹션", "핵심"]
    maxp = max((len(ps) for ps in groups.values()), default=0)
    ws.cell(1, 1, "Mapping audit: how every source label was mapped to a canonical row (override > seed > exact > embed > llm > new)").font = HEADER_FONT
    for c, h in enumerate(fixed, 1):
        ws.cell(2, c, h)
    col = len(fixed) + 1
    for i in range(maxp):
        for j, h in enumerate(MAPPING_PER_PRODUCT):
            ws.cell(2, col + 4 * i + j, h)
    review_col = col + 4 * maxp
    ws.cell(2, review_col, "검토 필요")
    for c in range(1, review_col + 1):
        t = ws.cell(2, c)
        t.font, t.fill, t.alignment = HEADER_FONT, HEADER_FILL, Alignment(vertical="center", wrap_text=True)
    for major, rows in model.items():
        ps = groups[major]
        band = max(ws.max_row + 1, 3)  # one band per category naming its products in the per-product column slots
        ws.cell(band, 1, major).font = HEADER_FONT
        ws.cell(band, 2, "제품 →").font = HEADER_FONT
        for i, p in enumerate(ps):
            ws.cell(band, col + 4 * i, clean_text(f"{p.brand} {p.model_number}")).font = HEADER_FONT
        for c in range(1, review_col + 1):
            ws.cell(band, c).fill = PatternFill("solid", fgColor="F2F2F2")
        for row in rows:
            if row["id"] in HEADER_IDS_SET or not any(row.get("sources", [])) or row["section"] == "기본정보":
                continue
            r = max(ws.max_row + 1, 3)
            vals = [major, row["id"], row["key_en"], row["key_ko"], row["section"], "Y" if row["core"] else ""]
            for c, v in enumerate(vals, 1):
                _neutralize(ws.cell(r, c, clean_text(v)))
            unsure = False
            for i, srcs in enumerate(row["sources"]):
                if not srcs:
                    continue
                cells = [NL.join(str(s["label"]) for s in srcs), NL.join(str(s["value"])[:200] for s in srcs),
                         NL.join(s["method"] for s in srcs), NL.join(f"{s['score']:.2f}" for s in srcs)]
                for j, v in enumerate(cells):
                    c = ws.cell(r, col + 4 * i + j, clean_text(v))
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    _neutralize(c)
                unsure = unsure or any(canon.needs_review(s["method"], s["score"]) for s in srcs)
            flag = ws.cell(r, review_col, "검토" if unsure else "")
            if unsure:
                flag.fill = UNSURE_FILL
    widths = [12, 26, 26, 22, 14, 6] + [30, 28, 9, 7] * maxp + [9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(review_col)}{max(ws.max_row, 3)}"


def _atomic_save(wb: Workbook, path: Path) -> Path:
    """Save to a temp file then os.replace; if the target is locked (open in Excel), use a timestamped name."""
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        os.replace(tmp, alt)
        print(f"[excel] {path} is locked; saved to {alt} instead", file=sys.stderr)
        return alt
