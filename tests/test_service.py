"""Plain-assert tests (pytest not installed). Run: python tests/test_service.py"""
import os
import sys
import tempfile
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

import catalog
import service
from catalog import Candidate
from schema import ProductRecord
from store import Store


class FakeAdapter:
    def __init__(self, fail_urls=()):
        self.scrape_calls = []
        self.fail_urls = set(fail_urls)

    def discover(self, category, limit=30):
        return [Candidate(brand="Fake", model_number=f"M{i}", name=f"Fridge {i}", url=f"http://f/{i}",
                          price_usd=None if i == 9 else 100.0 * (i + 1)) for i in range(10)][:limit]

    def scrape(self, url):
        self.scrape_calls.append(url)
        if url in self.fail_urls:
            raise RuntimeError("boom")
        n = url.rsplit("/", 1)[1]
        return (ProductRecord(brand="Fake", model_number=f"M{n}", product_name=f"Fridge {n}", product_url=url,
                              pod_features=["Ice maker", "Wi-Fi"]), [], [])


class patched:
    """Context manager: route catalog.adapter to a fake (no dependence on real adapter modules)."""
    def __init__(self, fake):
        self.fake = fake

    def __enter__(self):
        self.orig = catalog.adapter
        catalog.adapter = lambda brand: self.fake
        return self.fake

    def __exit__(self, *a):
        catalog.adapter = self.orig


def _cands(prices):
    return [Candidate(brand="B", model_number=f"M{i}", name="n", url=f"u{i}", price_usd=p)
            for i, p in enumerate(prices)]


def test_preset_terciles_and_unknown():
    bands = service.classify_bands(_cands([100, 200, 300, 400, 500, 600, None]), "preset")
    assert list(bands) == ["Budget", "Mid", "Premium", service.UNKNOWN_BAND]
    assert [c.price_usd for c in bands["Budget"]] == [100, 200]
    assert [c.price_usd for c in bands["Mid"]] == [300, 400]
    assert [c.price_usd for c in bands["Premium"]] == [500, 600]
    assert len(bands[service.UNKNOWN_BAND]) == 1


def test_preset_edge_cases():
    assert service.classify_bands([], "preset") == {b: [] for b in service.PRESET_BANDS}
    bands = service.classify_bands(_cands([None, None]), "preset")
    assert len(bands[service.UNKNOWN_BAND]) == 2 and "Budget" in bands


def test_custom_thresholds():
    bands = service.classify_bands(_cands([500, 1000, 1500, 3000, None]), "custom", [1000, 2000])
    assert list(bands) == ["< $1,000", "$1,000 - $2,000", ">= $2,000", service.UNKNOWN_BAND]
    assert [c.price_usd for c in bands["< $1,000"]] == [500]
    assert [c.price_usd for c in bands["$1,000 - $2,000"]] == [1000, 1500]  # boundary goes up
    assert [c.price_usd for c in bands[">= $2,000"]] == [3000]
    assert service.classify_bands(_cands([1, None]), "custom", [])["All prices"][0].price_usd == 1


def test_default_selection_caps():
    bands = service.classify_bands(_cands(list(range(100, 1900, 100))), "preset")
    assert len(service.default_selection(bands, 2)) == 6
    assert len(service.default_selection(bands, 10)) == service.MAX_SELECTED


def test_search_isolates_brand_failure_and_caches():
    fake = FakeAdapter()
    with tempfile.TemporaryDirectory() as t, patched(fake):
        s = Store(Path(t) / "c.db")
        orig = catalog.adapter
        catalog.adapter = lambda b: (_ for _ in ()).throw(KeyError(b)) if b == "Bad" else fake
        try:
            cands, log = service.search(["Fake", "Bad"], "refrigerator", 5, store=s)
            assert len(cands) == 5 and [l[1] for l in log] == ["ok", "failed"]
            _, log2 = service.search(["Fake"], "refrigerator", 5, store=s)
            assert "(cache)" in log2[0][2]
        finally:
            catalog.adapter = orig


def test_collect_isolates_failure_and_uses_cache():
    fake = FakeAdapter(fail_urls={"http://f/1"})
    cands = fake.discover("refrigerator", 3)
    events, sleeps = [], []
    with tempfile.TemporaryDirectory() as t, patched(fake):
        s = Store(Path(t) / "c.db")
        products, docs, specs, modes, log = service.collect(
            cands, progress_cb=lambda d, n, m: events.append((d, n)), store=s, delay=0.5, sleep=sleeps.append)
        assert [p.model_number for p in products] == ["M0", "M2"]
        assert [l[1] for l in log] == ["ok", "failed", "ok"]
        assert events[-1] == (3, 3) and len(sleeps) == 2
        assert products[0].price_usd == 100.0  # filled from candidate price
        # second run: successes come from cache, only the failing URL is re-scraped
        fake.scrape_calls.clear()
        products2, *_ = service.collect(cands, store=s, sleep=sleeps.append)
        assert fake.scrape_calls == ["http://f/1"] and len(products2) == 2


def test_collect_cancel():
    fake = FakeAdapter()
    cands = fake.discover("refrigerator", 4)
    cancel = threading.Event()

    def cb(done, total, msg):
        if done == 1:
            cancel.set()
    with patched(fake):
        products, *_, log = service.collect(cands, progress_cb=cb, cancel_event=cancel, sleep=lambda s: None)
    assert len(products) == 1 and log[-1][1] == "cancelled" and len(fake.scrape_calls) == 1


def test_export_excel_has_pod_compare():
    fake = FakeAdapter()
    with tempfile.TemporaryDirectory() as t, patched(fake):
        products, docs, specs, modes, log = service.collect(fake.discover("refrigerator", 2), sleep=lambda s: None)
        out = service.export_excel(products, docs, specs, modes, log, Path(t) / "o.xlsx")
        names = load_workbook(out).sheetnames
        for n in ("Products", "Run_Log", "POD_Items", "POD_Compare"):
            assert n in names, names


def test_load_config_registers_brands():
    cfg = service.load_config()
    assert [b["name"] for b in cfg["brands"]][:6] == ["Samsung", "LG", "KitchenAid", "GE", "Whirlpool", "Bosch"]
    assert len(cfg["brands"]) == 30 and cfg["brands"][-1]["name"] == "Panasonic"  # 6 original + 24 new (brands.yaml)
    assert catalog.ADAPTERS["LG"] == "lg_us" and catalog.ADAPTERS["Bosch"] == "bosch_us"
    assert [c["key"] for c in cfg["categories"] if c["enabled"]] == ["refrigerator"]


def test_app_smoke_search_flow():
    from streamlit.testing.v1 import AppTest
    fake = FakeAdapter()
    with tempfile.TemporaryDirectory() as t, patched(fake):
        os.environ["FRIDGE_DB"] = str(Path(t) / "c.db")
        service.load_config()
        at = AppTest.from_file(str(Path(__file__).resolve().parent.parent / "app.py"), default_timeout=30).run()
        assert not at.exception, at.exception
        at.multiselect[0].set_value(["Samsung"]).run()  # adapter patched: any brand hits the fake
        search_btn = next(b for b in at.button if "Search" in b.label)
        search_btn.click().run()
        assert not at.exception, at.exception
        assert len(at.session_state["candidates"]) == 10
        boxes = [c for c in at.checkbox if str(c.key).startswith("sel_")]
        assert len(boxes) == 10 and sum(1 for c in boxes if c.value) <= service.MAX_SELECTED
        os.environ.pop("FRIDGE_DB")


def test_polite_delay_applies_after_a_failed_scrape():
    fake = FakeAdapter(fail_urls={"http://f/0"})
    cands = fake.discover("refrigerator", 3)
    sleeps = []
    with patched(fake):
        products, *_, log = service.collect(cands, delay=0.5, sleep=sleeps.append)
    assert [l[1] for l in log] == ["failed", "ok", "ok"]
    assert len(sleeps) == 2  # a failed request still counts as hitting the site


def test_collect_failure_does_not_print_traceback_or_leak_detail(capsys=None):
    import contextlib
    import io
    fake = FakeAdapter(fail_urls={"http://f/0"})
    out = io.StringIO()
    with patched(fake), contextlib.redirect_stdout(out):
        *_, log = service.collect(fake.discover("refrigerator", 1), sleep=lambda s: None)
    assert "Traceback" not in out.getvalue()
    assert log[0][1] == "failed" and "RuntimeError" in log[0][2] and "boom" not in log[0][2]


def test_search_cache_is_per_limit_and_skips_empty():
    fake = FakeAdapter()
    with tempfile.TemporaryDirectory() as t, patched(fake):
        s = Store(Path(t) / "c.db")
        service.search(["Fake"], "refrigerator", 5, store=s)
        _, log = service.search(["Fake"], "refrigerator", 10, store=s)
        assert "(cache)" not in log[0][2], "limit=10 must not be served from the limit=5 cache entry"
        assert len(service.search(["Fake"], "refrigerator", 10, store=s)[0]) == 10
        fake.discover = lambda category, limit=30: []
        service.search(["Empty"], "refrigerator", 5, store=s)
        _, log = service.search(["Empty"], "refrigerator", 5, store=s)
        assert "(cache)" not in log[0][2]


def test_modes_added_to_cached_product_keep_scrape_age():
    import modes as modes_mod
    from datetime import datetime
    from schema import ModeRecord
    fake = FakeAdapter()
    cand = fake.discover("refrigerator", 1)
    orig = modes_mod.extract_modes
    modes_mod.extract_modes = lambda product, docs: [
        ModeRecord(brand="Fake", model_number="M0", mode_name="Vac", description="d", source_doc="m.pdf")]
    try:
        with tempfile.TemporaryDirectory() as t, patched(fake):
            s = Store(Path(t) / "c.db")
            service.collect(cand, store=s, sleep=lambda x: None)
            t0 = s.get_product("http://f/0").fetched_at
            import time as _t
            _t.sleep(1.1)
            service.collect(cand, store=s, with_modes=True, sleep=lambda x: None)
            hit = s.get_product("http://f/0")
            assert hit.modes and hit.fetched_at == t0
    finally:
        modes_mod.extract_modes = orig


class SubAdapter:
    SUPPORTED_SUBCATEGORIES = {"front_load", "dryer", "induction"}

    def __init__(self):
        self.calls = []

    def discover(self, subcategory, limit=30):
        self.calls.append(subcategory)
        if subcategory not in self.SUPPORTED_SUBCATEGORIES:
            raise ValueError(subcategory)
        return [Candidate(brand="Sub", model_number=f"{subcategory}{i}", name="n",
                          url=f"https://s/{subcategory}/{i}", price_usd=100.0 + i) for i in range(3)][:limit]

    def scrape(self, url):
        return ProductRecord(brand="Sub", model_number=url.rsplit("/", 1)[1], product_name="p", product_url=url), [], []


def test_search_subcategories_skips_unsupported_and_tags_candidates():
    ad = SubAdapter()
    with patched(ad):
        cands, log = service.search(["Sub"], ["front_load", "induction", "radiant", "top_load"], 5)
    assert ad.calls == ["front_load", "induction"]  # unsupported combos never reach the adapter
    assert ("Sub", "skipped", "라디언트 미지원") in log and ("Sub", "skipped", "전자동/탑로더 미지원") in log
    assert log[-1][1] == "ok" and log[-1][2] == "6 candidates"
    assert {(c.category, c.subcategory) for c in cands} == {("washer", "front_load"), ("cooking", "induction")}


def test_search_legacy_major_means_all_supported_subs_and_dedupes():
    class Fridge:
        def __init__(self):
            self.calls = []

        def discover(self, sub, limit=30):
            self.calls.append(sub)
            return [Candidate(brand="F", model_number="M", name="n", url="https://f/1", price_usd=1.0)]
    f = Fridge()
    with patched(f):
        cands, log = service.search(["F"], "refrigerator", 5)
    assert f.calls == ["french_door", "side_by_side", "top_freezer", "bottom_freezer", "built_in"]  # default support
    assert len(cands) == 1 and cands[0].subcategory == "french_door" and [l[1] for l in log] == ["ok"]


def test_search_failure_of_one_sub_is_partial_and_cache_is_per_sub():
    ad = SubAdapter()
    orig = ad.discover
    ad.discover = lambda sub, limit=30: (_ for _ in ()).throw(RuntimeError("x")) if sub == "dryer" else orig(sub, limit)
    with tempfile.TemporaryDirectory() as t, patched(ad):
        s = Store(Path(t) / "c.db")
        cands, log = service.search(["Sub"], ["front_load", "dryer"], 5, store=s)
        assert log[-1][1] == "partial" and len(cands) == 3
        ad.calls.clear()
        service.search(["Sub"], ["front_load", "induction"], 5, store=s)
        assert ad.calls == ["induction"]  # front_load served from its own cache entry


def test_group_by_major_and_collect_fills_category():
    ad = SubAdapter()
    with patched(ad):
        cands, _ = service.search(["Sub"], ["front_load", "induction"], 5)
        groups = service.group_by_major(cands)
        assert list(groups) == ["washer", "cooking"]
        products, *_ = service.collect(cands[:1] + cands[3:4], sleep=lambda s: None)
    assert [(p.category, p.subcategory) for p in products] == [("washer", "front_load"), ("cooking", "induction")]


def test_collect_user_selected_sub_is_authoritative(caplog=None):
    import logging

    class Mismatch:
        def scrape(self, url):  # PDP says side_by_side / washer, user picked french_door from the listing
            return (ProductRecord(brand="Fake", model_number="M1", product_name="n", product_url=url,
                                  category="washer", subcategory="side_by_side"), [], [])

    records = []

    class H(logging.Handler):
        def emit(self, record):
            records.append(record)
    h = H(level=logging.WARNING)
    service.logger.addHandler(h)
    try:
        cand = Candidate(brand="Fake", model_number="M1", name="n", url="http://f/1", category="refrigerator",
                         subcategory="french_door")
        with patched(Mismatch()):
            products, *_ = service.collect([cand], sleep=lambda s: None)
    finally:
        service.logger.removeHandler(h)
    assert (products[0].subcategory, products[0].category) == ("french_door", "refrigerator")
    assert any("side_by_side" in r.getMessage() and "french_door" in r.getMessage() for r in records), records
    # agreeing adapter: no warning, legacy label preserved
    records.clear()
    service.logger.addHandler(h)
    try:
        class Same:
            def scrape(self, url):
                return (ProductRecord(brand="Fake", model_number="M1", product_name="n", product_url=url,
                                      category="Refrigerator", subcategory="french_door"), [], [])
        with patched(Same()):
            products, *_ = service.collect([cand], sleep=lambda s: None)
    finally:
        service.logger.removeHandler(h)
    assert not records and products[0].category == "Refrigerator"


def test_preset_bands_with_fewer_than_three_priced_items():
    assert service.preset_thresholds(_cands([None])) is None and service.preset_thresholds([]) is None
    assert service.preset_thresholds(_cands([100, 200, 300])) == [200, 300]
    b1 = service.classify_bands(_cands([500, None]), "preset")
    assert [len(b1[n]) for n in service.PRESET_BANDS] == [0, 1, 0] and len(b1[service.UNKNOWN_BAND]) == 1
    b2 = service.classify_bands(_cands([900, 500]), "preset")
    assert [[c.price_usd for c in b2[n]] for n in service.PRESET_BANDS] == [[500], [], [900]]
    assert service.preset_thresholds(_cands([900, 500])) == [900, 900]
    given = service.classify_bands(_cands([100, 200, 300]), "preset", [150, 250])  # supplied cuts are honoured
    assert [[c.price_usd for c in given[n]] for n in service.PRESET_BANDS] == [[100], [200], [300]]


class FlagAdapter:
    """GE-like wall oven: Air fry only appears inside the 'Oven Cooking Modes' row; curated key says '–'."""
    def __init__(self, image_url="https://images.salsify.com/x.jpg"):
        self.image_url = image_url

    def scrape(self, url):
        p = ProductRecord(brand="GE", model_number="PTS9200SNSS", product_name="Double wall oven", product_url=url,
                          category="cooking", subcategory="electric_oven", image_url=self.image_url,
                          extra_specs={"Air fry": "–", "Oven Cooking Modes": "Convection Bake | No Preheat Air Fry"})
        from schema import RawSpec
        return p, [], [RawSpec(brand="GE", model_number="PTS9200SNSS", source="web", key="Cleaning", value="Self Clean")]


class patched_images:
    """Replace common.download_image with a recorder; result(brand, model, url) is the relative path or None."""
    def __init__(self, result, root):
        self.result, self.root, self.calls = result, root, []

    def __enter__(self):
        import common
        self.common, self.orig, self.orig_root = common, common.download_image, common.ROOT
        common.ROOT = self.root

        def fake(brand, model, url):
            self.calls.append((brand, model, url))
            r = self.result(brand, model, url) if callable(self.result) else self.result
            if isinstance(r, Exception):
                raise r
            if r:
                (self.root / r).parent.mkdir(parents=True, exist_ok=True)
                (self.root / r).write_bytes(b"img")
            return r
        common.download_image = fake
        return self

    def __exit__(self, *a):
        self.common.download_image, self.common.ROOT = self.orig, self.orig_root


def _oven_cand():
    return Candidate(brand="GE", model_number="PTS9200SNSS", name="n", url="https://www.geappliances.com/p/1",
                     category="cooking", subcategory="electric_oven")


def test_collect_derives_flags_from_full_spec_table():
    with patched(FlagAdapter()), patched_images(None, Path(tempfile.mkdtemp())):
        products, *_ = service.collect([_oven_cand()], sleep=lambda s: None)
    x = products[0].extra_specs
    assert x["Air fry"] == "Yes (No Preheat Air Fry)", x  # '–' curated value overwritten, evidence kept
    assert x["Self clean"].startswith("Yes (") and x["Convection"].startswith("Yes (")


def test_collect_downloads_image_sets_path_and_caches_it():
    root = Path(tempfile.mkdtemp())
    rel = "downloads/images/ge/PTS9200SNSS.jpg"
    with tempfile.TemporaryDirectory() as t, patched(FlagAdapter()), patched_images(rel, root) as im:
        s = Store(Path(t) / "c.db")
        products, *_ = service.collect([_oven_cand()], store=s, sleep=lambda x: None)
        assert products[0].image_path == rel and im.calls == [("GE", "PTS9200SNSS", "https://images.salsify.com/x.jpg")]
        hit = s.get_product("https://www.geappliances.com/p/1")
        assert hit.product.image_path == rel and hit.product.image_url and hit.product.extra_specs["Air fry"].startswith("Yes")
        products, *_ = service.collect([_oven_cand()], store=s, sleep=lambda x: None)  # cache hit, file present
        assert products[0].image_path == rel and len(im.calls) == 1
        (root / rel).unlink()  # file vanished: re-download and persist again
        products, *_ = service.collect([_oven_cand()], store=s, sleep=lambda x: None)
        assert len(im.calls) == 2 and (root / rel).is_file() and products[0].image_path == rel


def test_collect_image_failure_never_fails_the_product():
    for result in (None, RuntimeError("cdn down")):
        with patched(FlagAdapter()), patched_images(result, Path(tempfile.mkdtemp())) as im:
            products, _, _, _, log = service.collect([_oven_cand()], sleep=lambda x: None)
            assert len(products) == 1 and products[0].image_path is None and log[0][1] == "ok" and len(im.calls) == 1
    with patched(FlagAdapter(image_url=None)), patched_images("x.jpg", Path(tempfile.mkdtemp())) as im:
        service.collect([_oven_cand()], sleep=lambda x: None)
        assert im.calls == []  # nothing to download without an image_url


def test_collect_keeps_existing_image_path_when_file_present():
    root = Path(tempfile.mkdtemp())

    class Pre(FlagAdapter):
        def scrape(self, url):
            p, d, r = super().scrape(url)
            p.image_path = "downloads/images/ge/keep.jpg"
            (root / "downloads/images/ge").mkdir(parents=True, exist_ok=True)
            (root / p.image_path).write_bytes(b"x")
            return p, d, r
    with patched(Pre()), patched_images("other.jpg", root) as im:
        products, *_ = service.collect([_oven_cand()], sleep=lambda x: None)
    assert im.calls == [] and products[0].image_path == "downloads/images/ge/keep.jpg"


class ByCountry:
    """Patch catalog.adapter(brand, country='us') to per-country fakes (kr = KRW market)."""
    def __init__(self):
        self.calls = []
        outer = self

        class A:
            def __init__(self, cc):
                self.cc = cc
                self.SUPPORTED_SUBCATEGORIES = {"french_door"}

            def discover(self, sub, limit=30):
                outer.calls.append(("discover", self.cc, sub))
                if self.cc == "kr":
                    return [Candidate(brand="Fake", model_number=f"K{i}", name=f"냉장고 {i}", url=f"https://kr/{i}",
                                      price_usd=None, price_local=1_000_000.0 * (i + 1), currency="KRW") for i in range(3)]
                return [Candidate(brand="Fake", model_number=f"U{i}", name=f"Fridge {i}", url=f"https://us/{i}",
                                  price_usd=100.0 * (i + 1)) for i in range(3)]

            def scrape(self, url):
                outer.calls.append(("scrape", self.cc, url))
                n = url.rsplit("/", 1)[1]
                return ProductRecord(brand="Fake", model_number=f"X{n}", product_name="n", product_url=url), [], []
        self.A = A

    def __enter__(self):
        self.orig = catalog.adapter
        catalog.adapter = lambda brand, country="us": self.A(country)
        return self

    def __exit__(self, *a):
        catalog.adapter = self.orig


def test_search_across_countries_tags_region_country_and_caches_per_country():
    with tempfile.TemporaryDirectory() as d, ByCountry() as bc:
        store = Store(Path(d) / "c.db")
        found, log = service.search(["Fake"], ["french_door"], 30, store=store, countries=["us", "kr"])
        by = {c.country: c for c in found}
        assert len(found) == 6 and by["us"].region == "na" and by["us"].currency == "USD"
        k = [c for c in found if c.country == "kr"][0]
        assert (k.region, k.currency, k.price_usd, k.price_local) == ("kr", "KRW", None, 1_000_000.0)
        assert ("discover", "kr", "french_door") in bc.calls and ("discover", "us", "french_door") in bc.calls
        n = len(bc.calls)
        found2, _ = service.search(["Fake"], ["french_door"], 30, store=store, countries=["us", "kr"])
        assert len(bc.calls) == n and len(found2) == 6  # both countries served from separate cache entries
        assert {c.country for c in found2} == {"us", "kr"}
        one, _ = service.search(["Fake"], ["french_door"], 30, store=store)  # default: us only
        assert {c.country for c in one} == {"us"}


def test_search_country_without_adapter_is_skipped_with_log():
    with ByCountry():
        found, log = service.search(["Fake"], ["french_door"], 30, countries=["us", "kr"])
        assert len(found) == 6
        found, log = service.search(["Fake"], ["side_by_side"], 30, countries=["us", "kr"])
        assert found == [] and any(r[1] == "skipped" for r in log)


def test_bands_use_local_price_when_no_usd_price():
    cs = [Candidate(brand="B", model_number=f"M{i}", name="n", url=f"u{i}", region="kr", country="kr", currency="KRW",
                    price_local=p) for i, p in enumerate([1e6, 2e6, 3e6, 4e6, 5e6, 6e6, None])]
    assert service.preset_thresholds(cs) == [3e6, 5e6]
    bands = service.classify_bands(cs, "preset")
    assert [c.price_local for c in bands["Budget"]] == [1e6, 2e6] and len(bands[service.UNKNOWN_BAND]) == 1
    custom = service.classify_bands(cs, "custom", [2.5e6, 4.5e6], currency="KRW")
    assert list(custom)[0].startswith("< KRW ") and len(custom["< KRW 2,500,000"]) == 2


def test_collect_uses_country_adapter_and_stamps_product_region():
    with ByCountry() as bc:
        cand = Candidate(brand="Fake", model_number="K1", name="n", url="https://kr/1", category="refrigerator",
                         subcategory="french_door", region="kr", country="kr", currency="KRW", price_local=2e6)
        products, *_ = service.collect([cand], sleep=lambda x: None)
        assert ("scrape", "kr", "https://kr/1") in bc.calls
        p = products[0]
        assert (p.region, p.country, p.currency, p.price_local, p.price_usd) == ("kr", "kr", "KRW", 2e6, None)


def test_max_collect_constant():
    assert service.MAX_COLLECT == 12 and service.MAX_COLLECT >= service.MAX_SELECTED


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
