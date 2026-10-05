"""Plain-assert tests. Run: python tests/test_modes.py"""
import os
import sys
import tempfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_writer
import llm
import modes
from openpyxl import load_workbook
from schema import DocumentRecord, ModeRecord, ProductRecord


def _prod():
    return ProductRecord(brand="GE", model_number="M1", product_name="x", product_url="http://x")


def test_schema():
    p = _prod()
    assert p.fridge_temp_range_f is None and p.freezer_temp_range_f is None
    m = ModeRecord(brand="GE", model_number="M1", mode_name="Sabbath Mode", category="Religious",
                   description="d", source_doc="a.pdf")
    assert m.source_page is None and m.setting_range is None and m.how_to_activate is None


def test_modes_sheet_and_schema_rows():
    m = ModeRecord(brand="GE", model_number="M1", mode_name="Vacation", category="Away",
                   description="d", setting_range="34-44 F", source_doc="a.pdf", source_page=7)
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        excel_writer.write_excel([_prod()], [], [], out, modes=[m])
        wb = load_workbook(out)
        excel_writer.write_excel([], [], [], Path(d) / "e.xlsx")  # backward compat, no modes
    ws = wb["Modes"]
    assert [c.value for c in ws[1]] == list(ModeRecord.model_fields) and ws["A1"].font.bold
    assert ws.max_row == 2 and ws.cell(2, 9).value == 7
    names = [r[0].value for r in wb["Schema"].iter_rows(min_row=2)]
    assert "fridge_temp_range_f" in names and "mode_name" in names


def test_formula_escape():
    p = ProductRecord(brand="=cmd()", model_number="+1", product_name="@x", product_url="-y")
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        excel_writer.write_excel([p], [], [], out)
        ws = load_workbook(out)["Products"]
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=cmd()"
    assert ws["B2"].data_type == "s" and ws["C2"].data_type == "s"


def test_atomic_save_locked_target():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        out.write_bytes(b"old")
        real = os.replace
        calls = []

        def fake(src, dst):
            calls.append(str(dst))
            if str(dst) == str(out):
                raise PermissionError("locked")
            return real(src, dst)

        with mock.patch("excel_writer.os.replace", fake):
            got = excel_writer.write_excel([], [], [], out)
        assert got != out and got.exists() and got.name.startswith("r_")
        assert out.read_bytes() == b"old" and not list(Path(d).glob("*.tmp"))
        got2 = excel_writer.write_excel([], [], [], out)  # normal path replaces atomically
        assert got2 == out and load_workbook(out)


def test_page_scoring_and_garbled():
    pages = ["intro text about nothing at all here", "Vacation mode and Sabbath mode and SuperCool",
             "\u00e4\u00f6\u00fc\u00e9\u00e8 \u4e2d\u6587 mode mode mode 123 456 789 ###", "ice filter humidity"]
    assert modes.alpha_ratio(pages[2]) < modes.MIN_ALPHA_RATIO
    assert modes.score_page(pages[1]) > modes.score_page(pages[0]) == 0
    assert modes.pick_pages(pages) == [1, 3]
    assert modes.pick_pages(["mode"] * 10, top=6) == [0, 1, 2, 3, 4, 5]


def test_llm_parse():
    assert llm.parse_json('<think>x</think>```json\n[{"a": 1}]\n```') == [{"a": 1}]
    assert llm.parse_json('Sure! Here: {"a": 2} done') == {"a": 2}
    assert llm.parse_json("nope") is None


def test_hallucination_filter():
    with tempfile.TemporaryDirectory() as d:
        import fitz
        pdf = Path(d) / "M1_Manual.pdf"
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "Vacation mode turns off the ice maker. Sabbath mode disables lights. "
                                              "Fresh food range 34 to 44 F.")
        doc.save(pdf)
        doc.close()
        rec = DocumentRecord(brand="GE", model_number="M1", doc_type="Manual", source_url="u",
                             local_path=str(pdf), sha256="a", size_bytes=1)
        replies = [[
            {"mode_name": "Vacation mode", "category": "Away", "description": "Turns off ice.", "page": 1},
            {"mode_name": "Turbo Teleport", "category": "Cooling", "description": "Invented.", "page": 1},
            {"mode_name": "Sabbath mode", "category": "Bogus", "description": "Lights off.", "page": 5},
            {"mode_name": "sabbath MODE", "category": "Religious", "description": "Lights off.",
             "setting_range": "n/a", "page": 1},
        ], {"fridge_temp_range_f": "34-44 F", "freezer_temp_range_f": "-99-12 F"}]
        fake = mock.Mock(side_effect=lambda prompt: replies.pop(0))
        p = _prod()
        with mock.patch.object(llm, "chat_json", fake):
            out = modes.extract_modes(p, [rec])
    assert [(m.mode_name, m.category, m.source_page, m.setting_range) for m in out] == [
        ("Vacation mode", "Away", 1, None), ("sabbath MODE", "Religious", 1, None)]
    assert p.fridge_temp_range_f == "34-44 F" and p.freezer_temp_range_f is None


def test_llm_unreachable_returns_empty():
    with tempfile.TemporaryDirectory() as d:
        import fitz
        pdf = Path(d) / "M1_Manual.pdf"
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "Vacation mode text here")
        doc.save(pdf)
        doc.close()
        rec = DocumentRecord(brand="GE", model_number="M1", doc_type="Manual", source_url="u",
                             local_path=str(pdf), sha256="a", size_bytes=1)
        with mock.patch.object(llm, "chat_json", lambda prompt: None):
            assert modes.extract_modes(_prod(), [rec]) == []


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("PASS", name)
