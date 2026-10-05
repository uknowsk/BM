"""Plain-assert tests. Run: python tests/test_pod.py"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pod
from openpyxl import Workbook
from schema import ModeRecord, ProductRecord


def _p(brand, feats, **kw):
    return ProductRecord(brand=brand, model_number=brand + "1", product_name="x", product_url="http://x",
                         pod_features=feats, **kw)


def _tmp():
    return Path(tempfile.mkdtemp()) / "cache.json"


def _no_llm(prompt):
    raise AssertionError("LLM must not be called")


def test_keyword_mapping():
    p = _p("GE", ["VitaFresh humidity drawer keeps produce fresh", "Wi-Fi remote control via app",
                  "Sabbath Mode", "Fingerprint resistant stainless finish", "ENERGY STAR certified",
                  "LED lighting", "Door ajar alarm"])
    got: dict[str, set] = {}
    for i in pod.normalize_pod(p, llm_fn=_no_llm, cache_path=_tmp()):
        got.setdefault(i.raw_text, set()).add(i.taxonomy_key)
    assert got["Wi-Fi remote control via app"] == {"wifi_app"}
    assert got["Sabbath Mode"] == {"sabbath_mode"}
    assert got["Fingerprint resistant stainless finish"] == {"finish"}
    assert got["ENERGY STAR certified"] == {"energy_star"}
    assert got["LED lighting"] == {"led_light"}
    assert got["Door ajar alarm"] == {"door_alarm"}
    assert {"fresh_tech", "humidity_drawer"} <= got["VitaFresh humidity drawer keeps produce fresh"]


def test_no_duplicate_raw_key_pairs():
    p = _p("GE", ["Sabbath Mode", "sabbath mode", "LED lighting"], wifi_supported=True, ice_maker=True)
    items = pod.normalize_pod(p, llm_fn=_no_llm, cache_path=_tmp())
    pairs = [(i.raw_text.lower(), i.taxonomy_key) for i in items]
    assert len(pairs) == len(set(pairs))
    assert sum(i.taxonomy_key == "sabbath_mode" for i in items) == 1
    assert {"wifi_app", "ice_maker"} <= {i.taxonomy_key for i in items}  # field evidence added
    p2 = _p("GE", ["Wi-Fi connected"], wifi_supported=True)
    items2 = pod.normalize_pod(p2, llm_fn=_no_llm, cache_path=_tmp())
    assert [i.taxonomy_key for i in items2].count("wifi_app") == 1  # not duplicated


def test_llm_invalid_key_and_unreachable():
    p = _p("GE", ["Zorblax quantum thing"])
    items = pod.normalize_pod(p, llm_fn=lambda _: {"0": "made_up_key"}, cache_path=_tmp())
    assert items[0].taxonomy_key == "other" and items[0].matched_by == "none"
    items = pod.normalize_pod(p, llm_fn=lambda _: None, cache_path=_tmp())
    assert items[0].taxonomy_key == "other"
    items = pod.normalize_pod(p, llm_fn=lambda _: {"0": "led_light"}, cache_path=_tmp())
    assert items[0].taxonomy_key == "led_light" and items[0].matched_by == "llm"


def test_cache_hit_avoids_second_call():
    calls = []

    def fake(prompt):
        calls.append(prompt)
        return {"0": "led_light"}
    cp = _tmp()
    p = _p("GE", ["Zorblax quantum thing"])
    pod.normalize_pod(p, llm_fn=fake, cache_path=cp)
    items = pod.normalize_pod(p, llm_fn=fake, cache_path=cp)
    assert len(calls) == 1 and items[0].taxonomy_key == "led_light" and items[0].matched_by == "llm"


def test_batched_single_call():
    calls = []

    def fake(prompt):
        calls.append(prompt)
        return [{"index": 0, "key": "led_light"}, {"index": 1, "key": "shelves"}]
    p = _p("GE", ["Zorblax one", "Qwerty two"])
    items = pod.normalize_pod(p, llm_fn=fake, cache_path=_tmp())
    assert len(calls) == 1 and [i.taxonomy_key for i in items] == ["led_light", "shelves"]


def _two():
    a = _p("GE", ["Sabbath Mode", "LED lighting", "=evil formula LED"])
    b = _p("Bosch", ["Sabbath Mode", "Door ajar alarm"])
    return [a, b]


def test_compare_difference_logic():
    ps = _two()
    its = [pod.normalize_pod(p, llm_fn=_no_llm, cache_path=_tmp()) for p in ps]
    rows = {r["key"]: r for r in pod.compare_rows(ps, its)}
    assert rows["sabbath_mode"]["diff"] == "Both"
    assert rows["led_light"]["diff"] == "Only GE"
    assert rows["door_alarm"]["diff"] == "Only Bosch"
    assert "wifi_app" not in rows  # neither -> omitted


def test_sheet_structure():
    ps = _two()
    wb = Workbook()
    pod.add_pod_sheets(wb, ps, llm_fn=_no_llm)
    assert "POD_Items" in wb.sheetnames and "POD_Compare" in wb.sheetnames
    wi = wb["POD_Items"]
    assert wi.freeze_panes == "A2" and wi["A1"].font.bold and wi.auto_filter.ref
    assert wi.max_row == 1 + 3 + 2 and wi["I1"].value == "Mode ref"
    wc = wb["POD_Compare"]
    hdr = [c.value for c in wc[1]]
    assert hdr == ["Korean name", "English name", "Category", "GE present", "GE wording", "GE source",
                   "Bosch present", "Bosch wording", "Bosch source", "Difference"]
    order = [t[1] for t in pod.TAXONOMY]
    idx = [order.index(wc.cell(row=r, column=3).value) for r in range(2, wc.max_row + 1)]
    assert idx == sorted(idx)
    diffs = {wc.cell(row=r, column=2).value: wc.cell(row=r, column=10).value for r in range(2, wc.max_row + 1)}
    assert diffs["Sabbath mode"] == "Both" and diffs["LED lighting"] == "Only GE"
    for r in range(2, wc.max_row + 1):
        assert wc.cell(row=r, column=4).value in ("✓", "–")
        assert wc.cell(row=r, column=7).value in ("✓", "–")
        assert wc.cell(row=r, column=6).value in ("web", "spec field", "")
    for row in wi.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert c.data_type == "s"
    pod.add_pod_sheets(wb, ps, llm_fn=_no_llm)
    assert wb.sheetnames.count("POD_Items") == 1


def _mode(brand, name, page, doc="Manual.pdf", desc="d"):
    return ModeRecord(brand=brand, model_number=brand + "1", mode_name=name, description=desc,
                      source_doc=doc, source_page=page)


def test_bundle_splits_into_multiple_keys():
    raw = "External controls with actual temperature display, child lock and door alarm"
    items = pod.normalize_pod(_p("GE", [raw]), llm_fn=_no_llm, cache_path=_tmp())
    assert {i.taxonomy_key for i in items} == {"temp_display", "child_lock", "door_alarm"}
    assert all(i.raw_text == raw for i in items) and len(items) == 3
    wf = pod.normalize_pod(_p("GE", ["Advanced water filtration", "Cubed/crushed ice"]),
                           llm_fn=_no_llm, cache_path=_tmp())
    assert [(i.raw_text[:8], i.taxonomy_key) for i in wf] == [("Advanced", "water_filter"), ("Cubed/cr", "ice_type")]


def test_modes_cross_reflection_sabbath_both():
    ge = _p("GE", ["Enhanced Shabbos Mode"])
    bosch = _p("Bosch", ["LED Lighting"])
    modes = [_mode("Bosch", "Sabbath mode", 16, "Bosch1_Manual.pdf"), _mode("GE", "Sabbath Mode", 12)]
    its = [pod.normalize_pod(p, llm_fn=_no_llm, cache_path=_tmp(), modes=modes) for p in (ge, bosch)]
    b_sab = [i for i in its[1] if i.taxonomy_key == "sabbath_mode"]
    assert len(b_sab) == 1 and b_sab[0].matched_by == "mode"
    assert b_sab[0].raw_text == "Sabbath mode (Modes: Manual p16)"
    g_sab = [i for i in its[0] if i.taxonomy_key == "sabbath_mode"]
    assert len(g_sab) == 1 and g_sab[0].matched_by == "keyword"  # not duplicated
    rows = {r["key"]: r for r in pod.compare_rows([ge, bosch], its)}
    r = rows["sabbath_mode"]
    assert r["diff"] == "Both" and r["source"] == ["both", "modes"]
    assert r["wording"][0] == "Enhanced Shabbos Mode | Modes: Sabbath Mode p12"
    assert rows["led_light"]["source"] == ["", "web"]


def test_modes_other_products_ignored_and_unknown_to_other():
    modes = [_mode("Bosch", "Sabbath mode", 16), _mode("Bosch", "Zorblax gizmo", 3, desc="weird")]
    its = pod.normalize_pod(_p("GE", ["LED lighting"]), llm_fn=_no_llm, cache_path=_tmp(), modes=modes)
    assert [i.taxonomy_key for i in its] == ["led_light"]
    its = pod.normalize_pod(_p("Bosch", []), llm_fn=lambda _: None, cache_path=_tmp(), modes=modes)
    assert {(i.taxonomy_key, i.matched_by) for i in its} == {("sabbath_mode", "mode"), ("other", "mode")}


def test_mode_keyword_names_and_llm_fallback():
    names = {"Turbo Cool": "super_cool_freeze", "Super freezing": "super_cool_freeze",
             "Vacation mode": "vacation_mode", "Energy-saving mode": "eco_energy_saving_mode",
             "Freshness mode": "freshness_mode", "Lock": "child_lock", "Door alarm": "door_alarm",
             "Filter change notification": "filter_reminder", "Interior lighting": "led_light",
             "LED Dispenser Light": "led_light", "Automatic Icemaker": "ice_maker"}
    for n, k in names.items():
        assert pod.keyword_keys(n) == [k], (n, pod.keyword_keys(n))
    modes = [_mode("GE", "Zorblax", 1)]
    its = pod.normalize_pod(_p("GE", []), llm_fn=lambda _: {"0": "shelves"}, cache_path=_tmp(), modes=modes)
    assert its[0].taxonomy_key == "shelves" and its[0].matched_by == "mode"


def test_multiple_modes_same_key_merge_and_backward_compat():
    modes = [_mode("Bosch", "Door alarm", 16), _mode("Bosch", "Alarm warning tone", 16)]
    its = pod.normalize_pod(_p("Bosch", []), llm_fn=_no_llm, cache_path=_tmp(), modes=modes)
    assert len(its) == 1 and its[0].mode_ref == "Modes: Alarm warning tone p16"
    base = pod.normalize_pod(_p("Bosch", ["Door ajar alarm"]), llm_fn=_no_llm, cache_path=_tmp())
    assert [i.mode_ref for i in base] == [None] and base[0].matched_by == "keyword"
    wb = Workbook()
    pod.add_pod_sheets(wb, _two(), llm_fn=_no_llm)  # modes omitted
    wb2 = Workbook()
    pod.add_pod_sheets(wb2, _two(), llm_fn=_no_llm, modes=[_mode("Bosch", "Vacation mode", 15)])
    wc = wb2["POD_Compare"]
    row = [r for r in range(2, wc.max_row + 1) if wc.cell(row=r, column=2).value == "Vacation / away mode"]
    assert len(row) == 1 and wc.cell(row=row[0], column=9).value == "modes"


def test_pod_put_neutralises_formula_prefixes():
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(["=1+1", "+1", "-1", "@x", "\t=1", "\r=1", " =1", "plain"], 1):
        c = pod._put(ws, i, 1, v)
        assert c.data_type == "s" and bool(c.quotePrefix) == (v != "plain"), repr(v)


def _cat(brand, feats, category, **kw):
    return ProductRecord(brand=brand, model_number=brand + "1", product_name="x", product_url="http://x",
                         pod_features=feats, category=category, **kw)


def test_washer_and_cooking_taxonomies():
    w = _cat("LG", ["Steam sanitize cycle", "Auto-dispense detergent", "Stainless steel drum", "Quick Wash",
                    "Vibration reduction", "Stackable"], "washer", wifi_supported=True, energy_star=True)
    keys = {i.taxonomy_key for i in pod.normalize_pod(w, llm_fn=_no_llm, cache_path=_tmp())}
    assert {"steam", "sanitize", "auto_dispense", "stainless_drum", "quick_wash", "vibration_noise",
            "stackable", "wifi_app", "energy_star"} <= keys
    c = _cat("GE", ["Air Fry", "True convection", "Self-clean", "Center oval griddle", "Meat probe",
                    "Sabbath mode", "21,000 BTU power boil burner"], "cooking")
    keys = {i.taxonomy_key for i in pod.normalize_pod(c, llm_fn=_no_llm, cache_path=_tmp())}
    assert {"air_fry", "convection", "self_clean", "griddle", "temp_probe", "sabbath_mode", "boost_burner"} <= keys
    s = _cat("Samsung", ["Microwave power 900W", "Speed cook combi modes", "Light wave oven", "Convection"], "cooking")
    keys = {i.taxonomy_key for i in pod.normalize_pod(s, llm_fn=_no_llm, cache_path=_tmp())}
    assert {"microwave_power", "speed_cook", "convection"} <= keys
    assert pod.keyword_keys("Steam", "washer") == ["steam"] and pod.keyword_keys("Steam clean", "cooking") == ["self_clean"]


def test_llm_is_scoped_to_product_taxonomy():
    seen = []

    def llm(prompt):
        seen.append(prompt)
        return {"0": "heat_pump"}
    w = _cat("LG", ["zzz unknown thing"], "washer")
    items = pod.normalize_pod(w, llm_fn=llm, cache_path=_tmp())
    assert items[0].taxonomy_key == "heat_pump" and "heat_pump" in seen[0] and "ice_maker" not in seen[0]
    f = _p("GE", ["zzz unknown thing"])  # fridge: heat_pump is not valid there
    assert pod.normalize_pod(f, llm_fn=llm, cache_path=_tmp())[0].taxonomy_key == "other"


def test_compare_refuses_to_mix_majors_and_sheets_split():
    ps = [_cat("LG", ["Steam"], "washer"), _p("GE", ["Wi-Fi app"]), _cat("Bosch", ["Convection"], "cooking")]
    its = [pod.normalize_pod(p, llm_fn=_no_llm, cache_path=_tmp()) for p in ps]
    try:
        pod.compare_rows(ps, its)
        raise AssertionError("expected ValueError")
    except ValueError:
        pass
    wb = Workbook()
    pod.add_pod_sheets(wb, ps, its)
    assert "POD_Compare" not in wb.sheetnames
    for n in ("POD_Compare_refrigerator", "POD_Compare_washer", "POD_Compare_cooking", "POD_Items"):
        assert n in wb.sheetnames, wb.sheetnames
    ws = wb["POD_Compare_washer"]
    assert [c.value for c in ws[1]][3] == "LG present" and ws.cell(2, 2).value == "Steam wash / steam care"
    wb2 = Workbook()
    pod.add_pod_sheets(wb2, ps[:1], its[:1])
    assert "POD_Compare" in wb2.sheetnames and "POD_Compare_washer" not in wb2.sheetnames


def test_major_of_product_falls_back_to_subcategory():
    mk = lambda **kw: ProductRecord(brand="X", model_number="M", product_name="n", product_url="http://x", **kw)
    assert pod.major_of_product(mk(category="Refrigerator")) == "refrigerator"
    assert pod.major_of_product(mk(category="Cooktop")) == "cooking"
    assert pod.major_of_product(mk(category="Dryer")) == "washer"
    assert pod.major_of_product(mk(category="mystery", subcategory="induction")) == "cooking"
    assert pod.major_of_product(mk(category="", subcategory="dryer")) == "washer"
    assert pod.major_of_product(mk(category="mystery", subcategory="nope")) == "refrigerator"


def test_put_strips_control_chars_and_caps_length():
    wb = Workbook()
    ws = wb.active
    c = pod._put(ws, 1, 1, "ab\x00c\x07d\x1fe")
    assert c.value == "abcde"
    assert len(pod._put(ws, 2, 1, "x" * 40000).value) == 32767
    assert pod._put(ws, 3, 1, "\x00=1+1").value == "=1+1" and ws.cell(3, 1).data_type == "s"  # still neutralised


def test_save_cache_is_atomic_and_leaves_no_tmp():
    import os
    d = Path(tempfile.mkdtemp())
    path = d / "cache.json"
    pod._save_cache(path, {"a": "1"})
    assert pod._load_cache(path) == {"a": "1"}
    orig = os.replace
    os.replace = lambda *a, **k: (_ for _ in ()).throw(OSError("boom"))
    try:
        pod._save_cache(path, {"a": "2"})  # failed swap: old file intact, no stray temp file
    finally:
        os.replace = orig
    assert pod._load_cache(path) == {"a": "1"} and [p.name for p in d.iterdir()] == ["cache.json"]


if __name__ == "__main__":
    n = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            n += 1
            print("ok", name)
    print(f"{n} tests passed")
