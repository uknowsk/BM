"""Plain-assert tests (pytest not installed). Run: python tests/test_excel.py"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

import excel_writer
from excel_writer import write_excel
from schema import DocumentRecord, ModeRecord, ProductRecord, RawSpec


def _data():
    products = [
        ProductRecord(brand="GE", model_number="GNE27", product_name="GE French Door", product_url="http://x/1",
                      pod_features=["Ice maker", "WiFi"], wifi_supported=True, width_in=35.75),
        ProductRecord(brand="Bosch", model_number="B36", product_name="Bosch 500", product_url="http://x/2"),
    ]
    docs = [DocumentRecord(brand="GE", model_number="GNE27", doc_type="Manual", source_url="http://x/m.pdf",
                           local_path="downloads/ge/GNE27_Manual.pdf", sha256="ab" * 32, size_bytes=10, pages=3)]
    raws = [RawSpec(brand="GE", model_number="GNE27", source="web", section="Dims", key="Width", value="35 3/4 in"),
            RawSpec(brand="Bosch", model_number="B36", source="x.pdf", key="Volts", value="115")]
    return products, docs, raws


def test_workbook_structure():
    products, docs, raws = _data()
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "sub" / "r.xlsx"
        write_excel(products, docs, raws, out)
        wb = load_workbook(out)
    assert wb.sheetnames == ["Compare", "Mapping", "Products", "Documents", "Raw_Specs", "Modes", "Category_Specs", "Schema"]

    ws = wb["Products"]
    header = [c.value for c in ws[1]]
    assert header == list(ProductRecord.model_fields) + ["image"]  # trailing picture column
    assert ws.max_row == 3
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws["A1"].font.bold
    row1 = dict(zip(header, [c.value for c in ws[2]]))
    assert row1["pod_features"] == "Ice maker; WiFi"
    assert row1["wifi_supported"] is True and row1["width_in"] == 35.75
    assert dict(zip(header, [c.value for c in ws[3]]))["price_usd"] is None

    ws = wb["Documents"]
    assert [c.value for c in ws[1]] == list(DocumentRecord.model_fields)
    assert ws.max_row == 2
    col = list(DocumentRecord.model_fields).index("local_path") + 1
    assert ws.cell(2, col).hyperlink is not None

    ws = wb["Raw_Specs"]
    assert [c.value for c in ws[1]] == list(RawSpec.model_fields)
    assert ws.max_row == 3

    ws = wb["Schema"]
    assert [c.value for c in ws[1]][:4] == ["field", "type", "description", "unit"]
    sheet_col = [c.value for c in ws[1]].index("sheet")
    got = [(r[sheet_col].value, r[0].value) for r in ws.iter_rows(min_row=2)]
    expected = [(sh, f) for sh, m in (("Products", ProductRecord), ("Documents", DocumentRecord),
                                      ("Raw_Specs", RawSpec),
                                      ("Modes", ModeRecord)) for f in m.model_fields]
    expected += [("Category_Specs", f) for f in excel_writer.CATEGORY_SPECS_HEADER]
    assert got == expected
    units = {r[0].value: r[3].value for r in ws.iter_rows(min_row=2)}
    assert units["width_in"] == "in" and units["energy_kwh_year"] == "kWh/yr"
    # every row has a Korean description
    kcol = [c.value for c in ws[1]].index("description_ko")
    assert all(r[kcol].value for r in ws.iter_rows(min_row=2))


def test_category_specs_sheet_and_extra_specs_column():
    products, docs, raws = _data()
    products.append(ProductRecord(brand="LG", model_number="WM4000", product_name="LG Washer", product_url="http://x/3",
                                  category="washer", subcategory="front_load",
                                  extra_specs={"Capacity (cu ft)": "4.5", "Spin speed (rpm)": "1300", "Steam": "=1+1"}))
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        write_excel(products, docs, raws, out)
        wb = load_workbook(out)
    ws = wb["Category_Specs"]
    assert [c.value for c in ws[1]] == ["brand", "model_number", "product_name", "category", "subcategory", "key", "value"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 3 and rows[1][:6] == ["LG", "WM4000", "LG Washer", "washer", "front_load", "Spin speed (rpm)"]
    assert rows[1][6] == "1300"
    steam = [r for r in ws.iter_rows(min_row=2) if r[5].value == "Steam"][0][6]
    assert steam.data_type == "s" and steam.quotePrefix  # formula-like text neutralised here too
    pw = wb["Products"]
    hdr = [c.value for c in pw[1]]
    assert pw.max_row == 4  # still one row per product
    row = dict(zip(hdr, [c.value for c in pw[4]]))
    assert row["subcategory"] == "front_load" and "Spin speed (rpm): 1300" in row["extra_specs"]
    sch = {r[0].value: r for r in wb["Schema"].iter_rows(min_row=2) if r[5].value == "Products"}
    assert "extra_specs" in sch and "subcategory" in sch and sch["extra_specs"][1].value.startswith("dict")


def test_run_log_and_empty():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "e.xlsx"
        write_excel([], [], [], out, run_log=[("GE", "ok", "1 product"), ("Bosch", "failed", "boom")])
        wb = load_workbook(out)
    assert wb["Products"].max_row == 1
    assert wb.sheetnames[-1] == "Run_Log"
    assert [c.value for c in wb["Run_Log"][3]] == ["Bosch", "failed", "boom"]


def test_formula_injection_neutralised():
    products, docs, raws = _data()
    evil = ["=1+1", "+1+1", "-2+3", "@SUM(1)", "\t=1+1", "\r=1", "  =1+1", "\n=1"]
    raws = [RawSpec(brand="GE", model_number="GNE27", source="web", key=f"k{i}", value=v) for i, v in enumerate(evil)]
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        write_excel(products, docs, raws, out)
        ws = load_workbook(out)["Raw_Specs"]
        vcol = [c.value for c in ws[1]].index("value")
        rows = list(ws.iter_rows(min_row=2))
        assert len(rows) == len(evil)
        for row in rows:
            assert row[vcol].data_type == "s" and row[vcol].quotePrefix, repr(row[vcol].value)


def test_hyperlinks_are_relative_without_username():
    import excel_writer
    products, docs, raws = _data()
    with tempfile.TemporaryDirectory() as d:
        saved = excel_writer.ROOT
        excel_writer.ROOT = Path(d)
        try:
            out = Path(d) / "output" / "r.xlsx"
            write_excel(products, docs, raws, out)
        finally:
            excel_writer.ROOT = saved
        ws = load_workbook(out)["Documents"]
        col = list(DocumentRecord.model_fields).index("local_path") + 1
        link = ws.cell(2, col).hyperlink.target
        assert link == "../downloads/ge/GNE27_Manual.pdf", link
        assert "file:" not in link and "Users" not in link


def test_illegal_control_chars_stripped_and_long_strings_capped():
    products, docs, raws = _data()
    products[0].product_name = "Fr\x00idge\x07\x1f ok"
    products[0].pod_features = ["x" * 40000]
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        write_excel(products, docs, raws, out)  # would raise IllegalCharacterError without cleaning
        ws = load_workbook(out)["Products"]
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["product_name"] == "Fridge ok" and len(row["pod_features"]) == 32767


def test_hyperlinks_only_for_paths_inside_downloads():
    products, docs, raws = _data()
    import copy
    outside = copy.deepcopy(docs[0])
    outside.local_path = "../secrets/passwords.pdf"
    outside2 = copy.deepcopy(docs[0])
    outside2.local_path = "output/other.pdf"
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        write_excel(products, [docs[0], outside, outside2], raws, out)
        ws = load_workbook(out)["Documents"]
    col = list(DocumentRecord.model_fields).index("local_path") + 1
    assert ws.cell(2, col).hyperlink is not None
    assert ws.cell(3, col).hyperlink is None and ws.cell(4, col).hyperlink is None
    assert ws.cell(3, col).value == "../secrets/passwords.pdf"  # text kept, just not linked


def _png(path: Path, size=(400, 300)):
    from PIL import Image
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (30, 90, 160)).save(path, "PNG")


def _with_root(root: Path, fn):
    saved = excel_writer.ROOT
    excel_writer.ROOT = root
    try:
        return fn()
    finally:
        excel_writer.ROOT = saved


def test_products_sheet_embeds_images_and_degrades_gracefully():
    products, docs, raws = _data()
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        _png(root / "downloads" / "images" / "ge" / "GNE27.png")
        products[0].image_path = "downloads/images/ge/GNE27.png"
        products[1].image_path = "downloads/images/bosch/missing.png"  # file gone: row stays, no picture
        out = root / "output" / "r.xlsx"
        _with_root(root, lambda: write_excel(products, docs, raws, out))
        wb = load_workbook(out)  # the workbook must still open
        ws = wb["Products"]
        assert len(ws._images) == 1
        anchor = ws._images[0].anchor._from
        assert (anchor.col, anchor.row) == (len(ProductRecord.model_fields), 1)  # 0-based: 'image' column, row 2
        assert ws.row_dimensions[2].height and ws.row_dimensions[2].height > 60
        assert ws.row_dimensions[3].height is None or ws.row_dimensions[3].height < 60
        assert ws.max_row == 3 and ws.freeze_panes == "A2"
        # a path outside downloads/ is never embedded
        products[0].image_path = "../secrets/x.png"
        _png(root.parent / "x_secret.png")
        out2 = root / "output" / "r2.xlsx"
        _with_root(root, lambda: write_excel(products, docs, raws, out2))
        assert len(load_workbook(out2)["Products"]._images) == 0


def test_images_skipped_when_pillow_unavailable():
    products, docs, raws = _data()
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        _png(root / "downloads" / "images" / "ge" / "GNE27.png")
        products[0].image_path = "downloads/images/ge/GNE27.png"
        saved = excel_writer._embed_images
        excel_writer._embed_images = lambda *a, **k: (_ for _ in ()).throw(ImportError("no PIL"))
        try:
            out = root / "r.xlsx"
            _with_root(root, lambda: write_excel(products, docs, raws, out))
        finally:
            excel_writer._embed_images = saved
        assert load_workbook(out)["Products"].max_row == 3


def _row_of(ws, key_en, col=3):
    return next(r for r in ws.iter_rows(min_row=7) if r[col - 1].value == key_en)


def test_compare_sheet_is_transposed_with_real_ovens():
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import test_ge_us
    import test_kitchenaid_us
    ge = test_ge_us._rec("wall_double")[0]
    ka = test_kitchenaid_us._rec("KOEC730SWH")[0]
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        _png(root / "downloads" / "images" / "ge" / "PTS9200SNSS.png", (400, 300))
        _png(root / "downloads" / "images" / "kitchenaid" / "KOEC730SWH.png", (300, 400))
        ge.image_path, ka.image_path = "downloads/images/ge/PTS9200SNSS.png", "downloads/images/kitchenaid/KOEC730SWH.png"
        out = root / "output" / "r.xlsx"
        _with_root(root, lambda: write_excel([ge, ka], [], [], out))
        wb = load_workbook(out)  # still opens
    assert wb.sheetnames[0] == "Compare" and "Products" in wb.sheetnames
    ws = wb["Compare"]
    assert len(ws._images) == 2  # one picture per product column in the header block
    assert [(i.anchor._from.col, i.anchor._from.row) for i in ws._images] == [(4, 0), (5, 0)]
    assert [c.value for c in ws[6]][:6] == ["구분", "항목", "Item (EN)", "단위", "GE PTS9200SNSS", "KitchenAid KOEC730SWH"]
    assert ws["E3"].value == "PTS9200SNSS" and ws["F3"].value == "KOEC730SWH" and ws["E2"].value == "GE"
    assert ws.freeze_panes == "E7" and ws.auto_filter.ref.startswith("A6")
    assert ws.page_setup.orientation == "landscape" and ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.column_dimensions["G"].hidden  # diff column (filtering) is hidden
    sections = [r[0].value for r in ws.iter_rows(min_row=7) if r[0].value]
    assert sections[0] == "기본정보" and {"치수·무게", "조리(오븐·쿡탑)", "연결성"} <= set(sections)
    # one product per column, one CANONICAL attribute per row: Air fry is a check for both, with each product's wording in the cell/comment
    air = _row_of(ws, "Air fry")
    assert air[0].value == "조리(오븐·쿡탑)" and air[4].value.startswith("✓") and air[5].value.startswith("✓")
    assert "No Preheat Air Fry" in air[4].value and air[4].fill.fgColor.rgb.endswith("C6EFCE")
    assert air[4].comment is not None and "원문" in air[4].comment.text
    # core rows are expanded (outline level 1), long-tail rows sit in a collapsed outline group (level 2, hidden) under a band
    width = _row_of(ws, "Width")
    assert width[4].value == 29.75 and width[5].value == 29.75 and width[3].value == "in"
    assert ws.row_dimensions[width[0].row].outline_level == 1 and not ws.row_dimensions[width[0].row].hidden
    assert "Overall Width" in width[4].comment.text and "Dimensions > Width" in width[5].comment.text
    item = _row_of(ws, "Precision Cooking Modes")
    assert item[4].value.startswith("✓") and item[5].value == "–" and item[5].fill.fgColor.rgb.endswith("EDEDED")
    assert item[0].fill.fgColor.rgb.endswith("FFC000") and item[6].value == "diff"  # differing row marked
    assert ws.row_dimensions[item[0].row].outline_level == 2 and ws.row_dimensions[item[0].row].hidden
    band = next(r for r in ws.iter_rows(min_row=7) if r[1].value and "롱테일" in str(r[1].value))
    assert ws.row_dimensions[band[0].row].collapsed and ws.row_dimensions[band[0].row].outline_level == 1
    assert not any(isinstance(c.value, str) and " | " in c.value for r in ws.iter_rows(min_row=7) for c in r[:7])
    assert ws.row_dimensions[1].height and ws.row_dimensions[1].height > 80
    # Mapping audit sheet: canonical id | labels | section | core | per product (source label, value, method, score) | review flag
    mp = wb["Mapping"]
    head = [c.value for c in mp[2]]
    assert head[:6] == ["구분(major)", "canonical id", "항목 (EN)", "항목 (KO)", "섹션", "핵심"] and head[6:10] == ["원문 항목명", "원문 값", "방법", "점수"]
    assert head[-1] == "검토 필요"
    row = next(r for r in mp.iter_rows(min_row=3) if r[1].value == "width")
    assert row[0].value == "cooking" and row[3].value == "폭" and row[5].value == "Y"
    assert "Overall Width" in row[6].value and "Dimensions > Width" in row[10].value and row[8].value.startswith("seed") and row[9].value.startswith("1.00")
    assert row[-1].value in (None, "")
    assert mp.auto_filter.ref and mp.freeze_panes == "C3"


def test_compare_sheets_per_major_and_formula_text_neutralised():
    products, docs, raws = _data()
    products[0].category = products[1].category = "cooking"
    products[0].extra_specs = {"Info > Note": "=1+1"}
    products.append(ProductRecord(brand="LG", model_number="WM1", product_name="W", product_url="http://x/3", category="washer",
                                  extra_specs={"Capacity": "4.5 cu ft"}))
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "r.xlsx"
        write_excel(products, docs, raws, out)
        wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Compare_washer", "Compare_cooking"]
    note = _row_of(wb["Compare_cooking"], "Note")[4]
    assert note.data_type == "s" and note.quotePrefix
    assert [c.value for c in wb["Compare_washer"][6]][4] == "LG WM1"
    cap = _row_of(wb["Compare_washer"], "Washer capacity")
    assert cap[4].value == 4.5 and cap[3].value == "cu ft"


def test_no_compare_sheet_without_products_and_images_degrade():
    with tempfile.TemporaryDirectory() as d:
        write_excel([], [], [], Path(d) / "e.xlsx")
        assert "Compare" not in load_workbook(Path(d) / "e.xlsx").sheetnames
        products, docs, raws = _data()
        products[0].image_path = "downloads/images/ge/none.png"  # missing file
        out = Path(d) / "r.xlsx"
        write_excel(products, docs, raws, out)
        assert len(load_workbook(out)["Compare"]._images) == 0


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("PASS", name)
