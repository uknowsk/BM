"""Plain-assert tests (also pytest-compatible). Run: python tests/test_server.py
Uses FastAPI TestClient in FRIDGE_MOCK=1 mode (no network, no LLM)."""
import io
import os
import re
import sys
import tempfile
import time
from pathlib import Path

os.environ["FRIDGE_MOCK"] = "1"
os.environ["FRIDGE_OUTPUT_DIR"] = tempfile.mkdtemp(prefix="gauge_out_")
os.environ["PYTHONIOENCODING"] = "utf-8"
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import server

ORIGIN = f"http://127.0.0.1:{server.PORT}"
# browser-like client: loopback host + same-origin header on every request (POSTs require Origin)
client = TestClient(server.app, base_url=ORIGIN, headers={"origin": ORIGIN})


def wait(job_id, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("done", "error", "cancelled"):
            return j
        time.sleep(0.15)
    raise AssertionError("job timeout")


TIERS = ("t1", "t2", "t3", "t4", "t5")
TIER_NAMES = ["보급", "중저가", "중가", "중고가", "프리미엄"]


def run_search(**kw):
    body = {"brands": ["Samsung", "LG", "KitchenAid"], "category": "refrigerator", "limit": 30,
            "band_mode": "auto", "thresholds": None, **kw}
    r = client.post("/api/search", json=body)
    assert r.status_code == 200, r.text
    return wait(r.json()["job_id"])


def test_brands_and_categories():
    brands = {b["name"]: b for b in client.get("/api/brands").json()}
    assert all(brands[n]["enabled"] for n in ("Samsung", "LG", "KitchenAid", "GE", "Whirlpool", "Bosch"))
    assert not brands["Miele"]["enabled"] and brands["Miele"]["note"] and not brands["AEG"]["enabled"]
    assert brands["Miele"]["categories"] == [] and brands["Miele"]["majors"] == []
    assert "induction" in brands["KitchenAid"]["categories"] and "front_load" not in brands["KitchenAid"]["categories"]
    assert brands["Bosch"]["majors"] == ["refrigerator", "washer", "cooking"]
    cats = {c["key"]: c for c in client.get("/api/categories").json()}
    assert list(cats) == ["refrigerator", "washer", "cooking"] and cats["washer"]["label_ko"] == "세탁기"
    kids = {c["key"]: c for m in cats.values() for c in m["children"]}
    assert len(kids) == len(server.catalog.sub_keys()) == 18 and "gas_cooktop" in kids and kids["induction"]["label_ko"] == "인덕션"
    assert "Samsung" in kids["induction"]["brands"] and "KitchenAid" not in kids["front_load"]["brands"]
    assert kids["compact"]["enabled"] and kids["compact"]["brands"] == ["Whirlpool", "Bosch"]
    assert all(c["enabled"] == bool(c["brands"]) for c in kids.values()) and cats["cooking"]["enabled"]


def test_brands_registry_30_with_group_and_countries():
    brands = client.get("/api/brands").json()
    names = [b["name"] for b in brands]
    assert len(names) == len(set(names)) == 30 and names[:6] == ["Samsung", "LG", "KitchenAid", "GE", "Whirlpool", "Bosch"]
    by = {b["name"]: b for b in brands}
    assert all(b["group"] for b in brands)
    assert by["Maytag"]["group"] == "Whirlpool Corp." and by["Café"]["group"] == "Haier·GE" and by["Miele"]["group"] == "프리미엄"
    assert by["Samsung"]["countries"] == ["us"] and by["Bosch"]["enabled"]  # mock: only the US adapters exist
    for n in ("Maytag", "Fisher & Paykel", "De Dietrich", "Panasonic"):  # new brands: no adapter file yet, no error
        assert not by[n]["enabled"] and by[n]["note"] == "준비 중" and by[n]["countries"] == [] and by[n]["categories"] == []
    r = client.post("/api/search", json={"brands": ["Maytag"], "category": "refrigerator"})
    assert r.status_code == 422
    assert server.MAX_SEARCH_COMBOS == 120 and server.MAX_BRANDS == 30
    for name, dom in server.BRAND_DOMAINS.items():  # every brand has a host allowed for its own https site
        assert server.host_allowed(name, f"https://www.{dom}/x"), name
        assert not server.host_allowed(name, f"http://www.{dom}/x") and not server.host_allowed(name, "https://evil.example/x")


def test_search_cap_uses_max_combos():
    old = server.MAX_SEARCH_COMBOS
    server.MAX_SEARCH_COMBOS = 3
    try:
        r = client.post("/api/search", json={"brands": ["Samsung", "LG"], "subcategories": ["french_door", "side_by_side"]})
        assert r.status_code == 422 and "3" in r.json()["detail"]
    finally:
        server.MAX_SEARCH_COMBOS = old
    assert len(set(client.get("/api/brands").json()[0]["countries"])) >= 1


def test_search_auto_bands():
    j = run_search()
    assert j["status"] == "done", j
    bands = j["result"]["bands"]
    assert set(bands) == {*TIERS, "unknown"}
    assert j["result"]["labels"] == dict(zip(TIERS, TIER_NAMES), unknown="Price unknown")
    assert sum(len(v) for v in bands.values()) == 28
    assert len(bands["unknown"]) == 1 and len(j["result"]["thresholds"]) == 4
    assert j["result"]["thresholds"] == sorted(j["result"]["thresholds"])
    for lo, hi in zip(TIERS, TIERS[1:]):  # tiers are ordered by price; a price equal to a cut sits in the higher tier
        if bands[lo] and bands[hi]:
            assert max(c["price_usd"] for c in bands[lo]) <= min(c["price_usd"] for c in bands[hi])
    for i, k in enumerate(TIERS):
        assert all((i == 0 or c["price_usd"] >= j["result"]["thresholds"][i - 1]) and (i == 4 or c["price_usd"] < j["result"]["thresholds"][i])
                   for c in bands[k]), k


def test_search_custom_and_validation():
    j = run_search(band_mode="custom", thresholds=[500, 1000, 1500, 2500], brands=["LG"])
    assert j["status"] == "done" and j["result"]["thresholds"] == [500, 1000, 1500, 2500]
    assert j["result"]["labels"] == {"t1": "< $500", "t2": "$500 - $1,000", "t3": "$1,000 - $1,500", "t4": "$1,500 - $2,500",
                                      "t5": ">= $2,500", "unknown": "Price unknown"}
    b = j["result"]["bands"]
    assert all(c["price_usd"] < 500 for c in b["t1"]) and all(c["price_usd"] >= 2500 for c in b["t5"])
    assert all(1000 <= c["price_usd"] < 1500 for c in b["t3"])
    bad = {"brands": ["LG"], "category": "refrigerator", "limit": 30, "band_mode": "custom", "thresholds": [3000, 100, 50, 10]}
    assert client.post("/api/search", json=bad).status_code == 422
    for t in (None, [], [1000, 2500], [100, 200, 300], [100, 200, 300, 400, 500], [100, 200, 200, 300], [-1, 100, 200, 300],
              [100, 200, 300, 1_000_001], [300, 200, 100, 50]):  # count != 4, equal / descending, negative, over the cap
        assert client.post("/api/search", json={**bad, "thresholds": t}).status_code == 422, t
    ok = client.post("/api/search", json={**bad, "thresholds": [0, 100, 200, 1_000_000]})  # limits themselves are valid
    assert ok.status_code == 200 and wait(ok.json()["job_id"])["status"] == "done"
    assert client.post("/api/search", json={**bad, "band_mode": "auto", "brands": ["Miele"]}).status_code == 422
    assert client.post("/api/search", json={**bad, "band_mode": "auto", "category": "vacuum"}).status_code == 422
    for subs in (["nope"], ["washer"], [""]):  # unknown key / a major key given as a sub key
        assert client.post("/api/search", json={**bad, "band_mode": "auto", "category": None, "subcategories": subs}).status_code == 422, subs
    # KitchenAid has no washers: nothing to run -> 422
    assert client.post("/api/search", json={"brands": ["KitchenAid"], "subcategories": ["front_load"]}).status_code == 422


def test_search_groups_per_major_with_independent_bands():
    j = run_search(brands=["Samsung", "KitchenAid"], category=None,
                   subcategories=["french_door", "front_load", "dryer", "induction", "radiant"], limit=30)
    assert j["status"] == "done", j
    res = j["result"]
    assert [g["category"] for g in res["groups"]] == ["refrigerator", "washer", "cooking"]
    assert res["bands"] is None and res["thresholds"] is None  # multi-major: legacy keys not populated
    log = " | ".join(j["log"])
    assert "Samsung: 라디언트 미지원" in log and "KitchenAid: 전자동" not in log
    assert "KitchenAid: 드럼 미지원" in log
    for g in res["groups"]:
        assert set(g["bands"]) == {*TIERS, "unknown"} and g["label_ko"] and set(g["labels"]) == set(g["bands"])
        allc = [c for k in g["bands"] for c in g["bands"][k]]
        assert allc and all(c["category"] == g["category"] and c["subcategory"] for c in allc)
        priced = [c["price_usd"] for c in allc if c["price_usd"] is not None]
        if g["thresholds"]:  # quintiles computed from THIS group's prices only
            assert g["thresholds"] == sorted(g["thresholds"]) and len(g["thresholds"]) <= 4
            assert all(min(priced) <= t <= max(priced) for t in g["thresholds"])
    washer = res["groups"][1]
    assert {c["subcategory"] for k in washer["bands"] for c in washer["bands"][k]} <= {"front_load", "dryer"}
    assert all(c["brand"] == "Samsung" for k in washer["bands"] for c in washer["bands"][k])  # KitchenAid unsupported
    assert res["total"] == sum(g["total"] for g in res["groups"])


def test_search_single_major_keeps_legacy_bands_and_groups():
    j = run_search(brands=["LG"], category=None, subcategories=["front_load"])
    res = j["result"]
    assert len(res["groups"]) == 1 and res["bands"] == res["groups"][0]["bands"] and len(res["thresholds"]) in (1, 2, 3, 4)


def test_search_cap_on_combos():
    r = client.post("/api/search", json={"brands": ["Samsung", "LG", "GE", "Whirlpool", "Bosch", "KitchenAid"],
                                         "subcategories": catalog_sub_keys() * 2 + ["french_door"] * 40})
    assert r.status_code == 422 and client.get("/api/brands").status_code == 200


def catalog_sub_keys():
    return server.catalog.sub_keys()


def _pick(bands, brand, idx=0):
    pool = [c for k in TIERS for c in bands[k] if c["brand"] == brand and c["model_number"] != "LRFOS3016S"]
    return pool[idx]


def test_full_flow_and_excel():
    bands = run_search()["result"]["bands"]
    picks = [_pick(bands, b) for b in ("Samsung", "LG", "KitchenAid")]
    r = client.post("/api/collect", json={"urls": picks, "with_modes": True, "browser_mode": "headless"})
    assert r.status_code == 200, r.text
    assert os.environ["FRIDGE_BROWSER_MODE"] == "headless" and os.environ["FRIDGE_HEADLESS"] == "1"
    jid = r.json()["job_id"]
    j = wait(jid)
    assert j["status"] == "done", j
    res = j["result"]
    assert len(res["products"]) == 3 and len(res["documents"]) == 6 and res["modes"]
    assert res["pod"]["rows"] and {"key", "ko", "present", "source", "diff"} <= set(res["pod"]["rows"][0])
    assert all(d["href"].startswith("/api/doc?path=downloads/") for d in res["documents"])
    assert [i["status"] for i in j["progress"]["items"]] == ["done"] * 3
    d = client.get(res["documents"][0]["href"])
    assert d.status_code == 200 and d.content.startswith(b"%PDF")

    x = client.get(f"/api/jobs/{jid}/excel")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = load_workbook(io.BytesIO(x.content))
    for sheet in ("Products", "Documents", "Modes", "POD_Items", "POD_Compare", "Run_Log"):
        assert sheet in wb.sheetnames, wb.sheetnames
    assert wb["Products"].max_row == 4


def test_collect_multi_model_per_brand_and_major_with_groups():
    j = run_search(brands=["Samsung", "LG"], category=None, subcategories=["french_door", "front_load", "induction"])
    picks = {}
    for g in j["result"]["groups"]:
        for k in TIERS:
            for c in g["bands"][k]:
                if c["model_number"] != "LRFOS3016S":
                    picks.setdefault((c["brand"], g["category"]), c)
    sel = [picks[("Samsung", "refrigerator")], picks[("Samsung", "washer")], picks[("LG", "cooking")]]
    r = client.post("/api/collect", json={"urls": sel, "with_modes": True})
    assert r.status_code == 200, r.text
    jid = r.json()["job_id"]
    res = wait(jid)["result"]
    assert [g["category"] for g in res["groups"]] == ["refrigerator", "washer", "cooking"]
    assert res["pod"]["rows"] == []  # mixed majors: no single matrix
    for g in res["groups"]:
        assert len(g["products"]) == 1 and g["pod"]["rows"] and g["products"][0]["category"] == g["category"]
    wash = res["groups"][1]["products"][0]
    assert wash["extra_specs"] and wash["capacity_fridge_cuft"] is None and wash["ice_maker"] is None
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{jid}/excel").content))
    for n in ("Category_Specs", "POD_Compare_refrigerator", "POD_Compare_washer", "POD_Compare_cooking"):
        assert n in wb.sheetnames, wb.sheetnames
    assert "POD_Compare" not in wb.sheetnames and wb["Category_Specs"].max_row > 1
    # several models of the same brand + major are accepted (no one-per-brand cap any more)
    two = [picks[("Samsung", "refrigerator")], {**picks[("Samsung", "refrigerator")], "url": picks[("Samsung", "refrigerator")]["url"] + "x"}]
    r = client.post("/api/collect", json={"urls": two})
    assert r.status_code == 200, r.text
    wait(r.json()["job_id"])


def test_sco_accepted_and_scr_rejected():
    ok = {"brands": ["Samsung"], "band_mode": "auto", "thresholds": None, "limit": 10}
    r = client.post("/api/search", json={**ok, "subcategories": ["sco"]})
    assert r.status_code == 200, r.text
    wait(r.json()["job_id"])
    assert client.post("/api/search", json={**ok, "subcategories": ["scr"]}).status_code == 422
    assert client.post("/api/search", json={**ok, "category": "cooking", "subcategories": ["scr"]}).status_code == 422
    kids = {c["key"]: c for m in client.get("/api/categories").json() for c in m["children"]}
    assert "scr" not in kids and kids["sco"]["label_ko"] == "SCO (스피드쿡 오븐)" and "Samsung" in kids["sco"]["brands"]


def test_collect_category_validation():
    ok = {"brand": "LG", "url": "https://www.lg.com/us/front-load/x/"}
    for bad in ({"subcategory": "nope"}, {"category": "vacuum"}, {"subcategory": "front_load", "category": "cooking"},
                {"subcategory": "radiant"}):  # LG mock does not support radiant
        assert client.post("/api/collect", json={"urls": [{**ok, **bad}]}).status_code == 422, bad
    two = [{**ok, "subcategory": "front_load"}, {**ok, "url": ok["url"] + "y", "subcategory": "dryer"}]
    r = client.post("/api/collect", json={"urls": two})
    assert r.status_code == 200, r.text  # same brand, same major, different subs: allowed
    wait(r.json()["job_id"])


def test_failure_is_isolated():
    bands = run_search(brands=["LG"])["result"]["bands"]
    allc = [c for k in bands for c in bands[k]]
    bad = next(c for c in allc if c["model_number"] == "LRFOS3016S")
    j = wait(client.post("/api/collect", json={"urls": [bad]}).json()["job_id"])
    assert j["status"] == "done" and j["result"]["products"] == []
    assert j["progress"]["items"][0]["status"] == "failed"
    assert any(r[1] == "failed" for r in j["result"]["run_log"])


def test_collect_validation():
    ok = {"brand": "LG", "url": "https://www.lg.com/us/refrigerators/lrfxc2416s/", "subcategory": "french_door"}
    for bad in (
        {"brand": "LG", "url": "http://www.lg.com/us/x"},                   # not https
        {"brand": "LG", "url": "https://evil.example.com/www.lg.com"},      # wrong domain
        {"brand": "LG", "url": "https://lg.com.evil.io/x"},                 # suffix trick
        {"brand": "Miele", "url": "https://www.miele.com/x"},               # unknown brand
        {"brand": "Samsung", "url": "https://www.lg.com/x"},                # domain of another brand
    ):
        assert client.post("/api/collect", json={"urls": [bad]}).status_code == 422, bad
    two = [ok, {**ok, "url": "https://www.lg.com/us/refrigerators/ltcs20020s/"}]
    r = client.post("/api/collect", json={"urls": two})
    assert r.status_code == 200, r.text  # two models of one brand in the same group
    wait(r.json()["job_id"])
    assert client.post("/api/collect", json={"urls": []}).status_code == 422
    assert client.post("/api/collect", json={"urls": [ok], "browser_mode": "x"}).status_code == 422


def test_doc_path_confined():
    for p in ("../server.py", "..%2Fserver.py", "downloads/../server.py", "C:/Windows/win.ini", "/etc/passwd"):
        assert client.get("/api/doc", params={"path": p}).status_code == 404, p
    assert client.get("/api/jobs/nope").status_code == 404
    assert client.get("/api/jobs/nope/excel").status_code == 404


def test_cancel():
    bands = run_search()["result"]["bands"]
    picks = [_pick(bands, b) for b in ("Samsung", "LG", "KitchenAid")]
    jid = client.post("/api/collect", json={"urls": picks}).json()["job_id"]
    assert client.post("/api/collect", json={"urls": picks[:1]}).status_code == 409  # one job at a time
    assert client.post(f"/api/jobs/{jid}/cancel").json()["ok"]
    j = wait(jid)
    assert j["status"] == "cancelled" and len(j["result"]["products"]) < 3


def test_static_and_headers():
    r = client.get("/")
    assert r.status_code == 200 and "Gauge" in r.text
    assert "frame-ancestors" in r.headers["content-security-policy"]
    assert client.get("/api/brands", headers={"host": "evil.com"}).status_code == 403
    assert client.post("/api/search", json={}, headers={"origin": "https://evil.com"}).status_code == 403
    fonts_css = client.get("/css/fonts.css")
    assert fonts_css.status_code == 200 and "googleapis" not in r.text and "gstatic" not in fonts_css.text
    woff = re.search(r"url\(\.\./(fonts/[^)]+\.woff2)\)", fonts_css.text).group(1)
    assert client.get("/" + woff).status_code == 200 and client.get("/js/theme.js").status_code == 200
    csp = r.headers["content-security-policy"]
    for d in ("base-uri 'none'", "form-action 'none'", "object-src 'none'", "default-src 'self'"):
        assert d in csp, csp
    assert "script-src" not in csp and "unsafe-inline" not in csp.split("style-src")[0]
    assert not re.findall(r"<script(?![^>]*\bsrc=)[^>]*>", r.text), "inline <script> would violate CSP"


def test_post_requires_exact_loopback_origin():
    bare = TestClient(server.app, base_url=ORIGIN)  # no default origin header
    assert bare.post("/api/search", json={}).status_code == 403  # missing Origin
    assert bare.get("/api/brands").status_code == 200  # GETs need none
    for o in ("null", "http://127.0.0.1:1", "http://localhost", "https://127.0.0.1:8765",
              "http://127.0.0.1:8765.evil.com", "http://evil.com", "http://localhost:8765/x"):
        assert bare.post("/api/search", json={}, headers={"origin": o}).status_code == 403, o
    for o in (f"http://127.0.0.1:{server.PORT}", f"http://localhost:{server.PORT}"):
        assert bare.post("/api/search", json={}, headers={"origin": o}).status_code == 422, o  # passed the guard
    h = {"origin": ORIGIN}
    assert bare.post("/api/search", json={}, headers={**h, "sec-fetch-site": "cross-site"}).status_code == 403
    assert bare.post("/api/search", json={}, headers={**h, "sec-fetch-site": "same-site"}).status_code == 403
    for v in ("same-origin", "none"):
        assert bare.post("/api/search", json={}, headers={**h, "sec-fetch-site": v}).status_code == 422, v


def test_testserver_host_gated_on_env():
    old = os.environ.pop("FRIDGE_TESTING", None)
    try:
        ts = TestClient(server.app)  # Host: testserver
        assert ts.get("/api/brands").status_code == 403
        os.environ["FRIDGE_TESTING"] = "1"
        assert ts.get("/api/brands").status_code == 200
    finally:
        os.environ.pop("FRIDGE_TESTING", None)
        if old is not None:
            os.environ["FRIDGE_TESTING"] = old


class _Adapter:
    def __init__(self, discover=None, scrape=None):
        self._d, self._s = discover, scrape

    def discover(self, category, limit=30):
        return self._d(category, limit)

    def scrape(self, url):
        return self._s(url)


class _patched_adapter:
    def __init__(self, adapter):
        self.adapter = adapter

    def __enter__(self):
        self.orig = server.catalog.adapter
        server.catalog.adapter = lambda brand: self.adapter

    def __exit__(self, *a):
        server.catalog.adapter = self.orig


def test_search_drops_candidates_with_disallowed_urls():
    C = server.Candidate
    bad = ["javascript:alert(1)", "http://www.lg.com/us/x", "https://evil.example.com/p", "https://lg.com.evil.io/p",
           "https://user:pw@www.lg.com/p"]
    cands = [C(brand="LG", model_number=f"B{i}", name="n", url=u, price_usd=100.0 + i) for i, u in enumerate(bad)]
    cands.append(C(brand="LG", model_number="GOOD", name="n", url="https://www.lg.com/us/refrigerators/good/", price_usd=500.0))
    with _patched_adapter(_Adapter(discover=lambda c, l: cands)):
        j = run_search(brands=["LG"])
    assert j["status"] == "done", j
    urls = [c["url"] for k in j["result"]["bands"] for c in j["result"]["bands"][k]]
    assert urls == ["https://www.lg.com/us/refrigerators/good/"] and j["result"]["total"] == 1


def _good_product(url, **kw):
    from schema import ProductRecord
    return ProductRecord(brand="LG", model_number="GOOD", product_name="Good", product_url=url, **kw)


def test_product_url_validated_before_rendering():
    url = "https://www.lg.com/us/refrigerators/good/"
    item = {"brand": "LG", "url": url, "model_number": "GOOD", "name": "n", "subcategory": "french_door"}
    with _patched_adapter(_Adapter(scrape=lambda u: (_good_product("javascript:alert(1)"), [], []))):
        j = wait(client.post("/api/collect", json={"urls": [item]}).json()["job_id"])
    assert j["status"] == "done" and j["result"]["products"][0]["product_url"] in (None, "")
    with _patched_adapter(_Adapter(scrape=lambda u: (_good_product(url), [], []))):
        j = wait(client.post("/api/collect", json={"urls": [item]}).json()["job_id"])
    assert j["result"]["products"][0]["product_url"] == url


def test_doc_only_pdf_and_sandboxed():
    d = server.DOWNLOADS_DIR / "mock"
    d.mkdir(parents=True, exist_ok=True)
    made = []
    try:
        for name, body in (("t_probe.pdf", b"%PDF-1.4\n"), ("t_probe.html", b"<script>1</script>"), ("t_probe.txt", b"x")):
            (d / name).write_bytes(body)
            made.append(d / name)
        ok = client.get("/api/doc", params={"path": "downloads/mock/t_probe.pdf"})
        assert ok.status_code == 200 and ok.headers["content-security-policy"].startswith("sandbox")
        assert ok.headers["x-content-type-options"] == "nosniff"
        for n in ("t_probe.html", "t_probe.txt"):
            assert client.get("/api/doc", params={"path": f"downloads/mock/{n}"}).status_code == 404, n
    finally:
        for f in made:
            f.unlink(missing_ok=True)


def test_errors_are_generic_and_logged_server_side():
    secret = r"C:\Users\victim\secret\token.txt"

    def boom(*a):
        raise RuntimeError(f"cannot open {secret}")
    with _patched_adapter(_Adapter(discover=boom, scrape=boom)):
        j = run_search(brands=["LG"])
        assert j["status"] == "error" and j["error"]
        item = {"brand": "LG", "url": "https://www.lg.com/us/refrigerators/x/", "model_number": "X", "name": "n",
                "subcategory": "french_door"}
        j2 = wait(client.post("/api/collect", json={"urls": [item]}).json()["job_id"])
    blob = repr(j) + repr(j2)
    assert "victim" not in blob and "secret" not in blob and "token.txt" not in blob, blob
    assert server._public_error(ValueError(secret)) == "ValueError"
    scrubbed = server._scrub(f"x {secret} y /home/victim/f.txt")
    assert "victim" not in scrubbed and "C:" not in scrubbed


def test_export_failure_keeps_results():
    url = "https://www.lg.com/us/refrigerators/good/"
    item = {"brand": "LG", "url": url, "model_number": "GOOD", "name": "n", "subcategory": "french_door"}
    orig = server._export
    server._export = lambda *a, **k: (_ for _ in ()).throw(OSError("disk full C:\\secret"))
    try:
        with _patched_adapter(_Adapter(scrape=lambda u: (_good_product(url), [], []))):
            j = wait(client.post("/api/collect", json={"urls": [item]}).json()["job_id"])
    finally:
        server._export = orig
    assert j["status"] == "done" and len(j["result"]["products"]) == 1 and not j["has_excel"]
    assert any("엑셀" in l and "OSError" in l for l in j["log"]) and "secret" not in repr(j)


def test_thread_start_failure_marks_job_error():
    class Boom:
        def __init__(self, *a, **k):
            pass

        def start(self):
            raise RuntimeError("cannot start new thread")
    orig = server.threading.Thread
    server.threading.Thread = Boom
    try:
        r = client.post("/api/search", json={"brands": ["LG"], "category": "refrigerator", "limit": 30, "band_mode": "auto"})
    finally:
        server.threading.Thread = orig
    assert r.status_code == 500
    assert not any(j.status in ("queued", "running") for j in server.JOBS.values())
    assert run_search(brands=["LG"])["status"] == "done"  # not locked out


def test_stale_job_does_not_lock_forever():
    stuck = server.Job("collect")
    stuck.status = "running"
    stuck.last_progress = time.time() - server.STALE_JOB_S - 1
    server.JOBS[stuck.id] = stuck
    r = client.post("/api/search", json={"brands": ["LG"], "category": "refrigerator", "limit": 30, "band_mode": "auto"})
    assert r.status_code == 200, r.text
    wait(r.json()["job_id"])
    assert stuck.status == "error" and stuck.cancel.is_set()
    fresh = server.Job("collect")
    fresh.status = "running"
    server.JOBS[fresh.id] = fresh
    try:
        assert client.post("/api/search", json={"brands": ["LG"], "category": "refrigerator", "limit": 30,
                                                "band_mode": "auto"}).status_code == 409
    finally:
        fresh.status = "done"


def test_browser_env_change_resets_adapter_mode_caches():
    import types
    calls = []
    mod = types.ModuleType("fake_adapter_mod")
    mod.reset_browser_mode = lambda: calls.append(1)
    plain = types.ModuleType("fake_adapter_plain")
    sys.modules["fake_adapter_mod"], sys.modules["fake_adapter_plain"] = mod, plain
    server.catalog.ADAPTERS["FakeA"], server.catalog.ADAPTERS["FakeB"] = "fake_adapter_mod", "fake_adapter_plain"
    try:
        server._set_browser_env("visible")
    finally:
        server.catalog.ADAPTERS.pop("FakeA"), server.catalog.ADAPTERS.pop("FakeB")
        sys.modules.pop("fake_adapter_mod"), sys.modules.pop("fake_adapter_plain")
        server._set_browser_env("auto")
    assert len(calls) == 1 and os.environ["FRIDGE_BROWSER_MODE"] == "auto"


def test_collect_requires_subcategory_and_never_trusts_category():
    base = {"brand": "LG", "url": "https://www.lg.com/us/x/a/"}
    for bad in ({}, {"category": "refrigerator"}, {"subcategory": ""}, {"subcategory": None}):  # omitted subcategory
        assert client.post("/api/collect", json={"urls": [{**base, **bad}]}).status_code == 422, bad
    # a forged category is rejected: the major is derived from the sub key
    a = {**base, "subcategory": "french_door"}
    b = {**base, "url": base["url"] + "b", "subcategory": "side_by_side", "category": "cooking"}
    assert client.post("/api/collect", json={"urls": [a, b]}).status_code == 422
    b_ok = {**b}
    b_ok.pop("category")
    r = client.post("/api/collect", json={"urls": [a, b_ok]})
    assert r.status_code == 200, r.text  # same brand + derived major twice: allowed now
    wait(r.json()["job_id"])
    # a client-supplied category is never used: a consistent one is accepted, product is tagged by sub
    r = client.post("/api/collect", json={"urls": [{**a, "category": "refrigerator"}]})
    assert r.status_code == 200, r.text
    wait(r.json()["job_id"])


def test_mock_js_only_served_in_mock_mode():
    html = client.get("/").text
    assert "js/mock.js" in html and client.get("/js/mock.js").status_code == 200
    old = os.environ["FRIDGE_MOCK"]
    os.environ["FRIDGE_MOCK"] = "0"
    try:
        assert client.get("/js/mock.js").status_code == 404
        assert "mock.js" not in client.get("/").text and "mock.js" not in client.get("/index.html").text
        assert client.get("/js/app.js").status_code == 200
    finally:
        os.environ["FRIDGE_MOCK"] = old
    static = (server.WEB_DIR / "index.html").read_text(encoding="utf-8")
    assert "mock.js" not in static  # never an unconditional tag in the source page


def test_expired_job_thread_cannot_overwrite_status_or_result():
    import threading
    entered, gate = threading.Event(), threading.Event()

    def slow_scrape(url):
        entered.set()
        gate.wait(15)
        return _good_product(url), [], []

    url = "https://www.lg.com/us/refrigerators/good/"
    item = {"brand": "LG", "url": url, "model_number": "GOOD", "name": "n", "subcategory": "french_door"}
    with _patched_adapter(_Adapter(scrape=slow_scrape)):
        jid = client.post("/api/collect", json={"urls": [item]}).json()["job_id"]
        job = server.JOBS[jid]
        assert entered.wait(10)
        job.last_progress = time.time() - server.STALE_JOB_S - 1
        # expiry happens on the next registration; the still-alive stale thread keeps the lock (documented choice)
        r = client.post("/api/search", json={"brands": ["LG"], "category": "refrigerator"})
        assert r.status_code == 409, r.text
        assert job.stale and job.status == "error" and job.cancel.is_set()
        err = job.error
        gate.set()
        job.thread.join(15)
        assert not job.thread.is_alive()
    assert job.status == "error" and job.error == err and job.result is None and job.xlsx_path is None
    snap = client.get(f"/api/jobs/{jid}").json()
    assert snap["status"] == "error" and snap["result"] is None
    assert run_search(brands=["LG"])["status"] == "done"  # thread gone: lock is free again


def test_finish_never_overwrites_terminal_or_stale():
    j = server.Job("collect")
    assert j.finish("error", error="x") and j.status == "error"
    assert not j.finish("done", result={"a": 1}) and j.status == "error" and j.result is None
    k = server.Job("collect")
    k.stale = True
    assert not k.finish("done", result={"a": 1}) and k.status == "queued" and k.result is None


def _search_prices(prices, **kw):
    C = server.Candidate
    cands = [C(brand="LG", model_number=f"P{i}", name="n", url=f"https://www.lg.com/us/p{i}/", price_usd=p)
             for i, p in enumerate(prices)]
    with _patched_adapter(_Adapter(discover=lambda c, l: cands)):
        j = run_search(brands=["LG"], subcategories=["french_door"], category=None, **kw)
    assert j["status"] == "done", j
    return j["result"]["groups"][0]


def test_preset_search_with_few_priced_items_has_fewer_bands():
    g = _search_prices((500.0, 900.0))  # 2 priced -> 2 tiers, 1 cut; only t1, t2 filled
    assert [len(g["bands"][k]) for k in TIERS] == [1, 1, 0, 0, 0]
    assert g["thresholds"] == [700.0] and g["labels"]["t1"] == "1/2단계" and g["labels"]["t2"] == "2/2단계"
    assert g["labels"]["t3"] == "중가" and g["labels"]["unknown"] == "Price unknown"
    g = _search_prices((500.0, 600.0, 700.0, 800.0))  # 4 priced -> 4 tiers
    assert [len(g["bands"][k]) for k in TIERS] == [1, 1, 1, 1, 0] and len(g["thresholds"]) == 3
    g = _search_prices((500.0, None))  # a lone priced item: no cuts, it is t1
    assert [len(g["bands"][k]) for k in TIERS] == [1, 0, 0, 0, 0] and len(g["bands"]["unknown"]) == 1 and g["thresholds"] is None
    g = _search_prices((None, None))
    assert sum(len(g["bands"][k]) for k in TIERS) == 0 and len(g["bands"]["unknown"]) == 2 and g["thresholds"] is None


def test_preset_search_five_or_more_priced_items_fill_all_tiers_and_ties_go_up():
    g = _search_prices((100.0, 200.0, 300.0, 400.0, 500.0, 600.0, None))
    assert g["thresholds"] == [200.0, 300.0, 400.0, 500.0]
    assert [[c["price_usd"] for c in g["bands"][k]] for k in TIERS] == [[100.0], [200.0], [300.0], [400.0], [500.0, 600.0]]
    assert g["labels"] == dict(zip(TIERS, TIER_NAMES), unknown="Price unknown")


def test_preset_thresholds_computed_once_by_service():
    calls = []
    orig = server.service.preset_thresholds
    server.service.preset_thresholds = lambda c: (calls.append(1), orig(c))[1]
    try:
        j = run_search(brands=["LG"], subcategories=["french_door"], category=None)
    finally:
        server.service.preset_thresholds = orig
    assert j["status"] == "done" and len(calls) == 1, calls


def test_img_endpoint_serves_only_files_under_downloads_images():
    from PIL import Image
    d = server.DOWNLOADS_DIR / "images" / "testbrand"
    d.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (40, 30), (10, 20, 30)).save(d / "T1.png")
    Image.new("RGB", (40, 30), (10, 20, 30)).save(d / "T2.jpg")
    (d / "notes.txt").write_text("secret")
    (server.DOWNLOADS_DIR / "images" / "top.png").write_bytes(b"x")
    try:
        r = client.get("/api/img/testbrand/T1.png")
        assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content[:4] == bytes([0x89, 0x50, 0x4E, 0x47])
        assert r.headers["x-content-type-options"] == "nosniff"
        assert "max-age=" in r.headers["cache-control"] and int(r.headers["cache-control"].split("max-age=")[1].split(",")[0]) >= 86400
        assert client.get("/api/img/testbrand/T2.jpg").headers["content-type"] == "image/jpeg"
        for bad in ("/api/img/testbrand/notes.txt", "/api/img/testbrand/missing.png", "/api/img/../doc/x.png",
                    "/api/img/testbrand/..%2F..%2Fserver.py", "/api/img/testbrand/%2e%2e%2fT1.png", "/api/img/..%2Ftestbrand/T1.png",
                    "/api/img/testbrand/T1.png%00.txt", "/api/img/testbrand%5C..%5CT1.png/x.png", "/api/img/top.png",
                    "/api/img/testbrand/T1.PNG.exe", "/api/img/.../T1.png", "/api/img/mock/..%5C..%5C..%5Cserver.py"):
            assert client.get(bad).status_code == 404, bad
    finally:
        for f in (d / "T1.png", d / "T2.jpg", d / "notes.txt", server.DOWNLOADS_DIR / "images" / "top.png"):
            f.unlink(missing_ok=True)
        d.rmdir()


def test_collect_results_carry_image_src_and_features():
    j = run_search(brands=["GE"], category=None, subcategories=["electric_oven"])
    cand = next(c for k in TIERS for c in j["result"]["groups"][0]["bands"][k])
    jid = client.post("/api/collect", json={"urls": [cand]}).json()["job_id"]
    j = wait(jid)
    assert j["status"] == "done", j
    p = j["result"]["products"][0]
    assert p["image_src"].startswith("/api/img/mock/") and j["result"]["groups"][0]["products"][0]["image_src"] == p["image_src"]
    img = client.get(p["image_src"])
    assert img.status_code == 200 and img.headers["content-type"].startswith("image/")
    x = p["extra_specs"]
    assert any(" > " in k for k in x)  # sectioned full-table keys
    assert x["Air fry"].startswith("Yes (")  # derived from the multi-value 'Oven Cooking Modes' row
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{jid}/excel").content))
    assert len(wb["Products"]._images) == 1 and wb.sheetnames[0] == "Compare" and len(wb["Compare"]._images) == 1
    cmp = j["result"]["groups"][0]["compare"]  # same rows the workbook's Compare sheet is built from
    assert {r["section"] for r in cmp} >= {"기본정보", "치수·무게", "조리(오븐·쿡탑)"} and all("core" in r and "sources" in r for r in cmp)
    assert next(r for r in cmp if r["id"] == "image")["values"][0] == p["image_src"]
    air = next(r for r in cmp if r["key_en"] == "Air fry")
    assert air["values"] == [True] and air["kind"] == "flag"
    assert any(r["id"] in ("bake", "broil", "convection-bake") for r in cmp), [r["id"] for r in cmp][:60]  # canonical rows (Bake, Broil, ...)
    assert wb["Compare"]["E3"].value == p["model_number"]


def test_product_without_image_has_null_image_src():
    p = server.ProductRecord(brand="GE", model_number="X", product_name="n", product_url="https://www.geappliances.com/x",
                             image_path="downloads/images/ge/none.jpg")
    assert server._product_dict(p)["image_src"] is None  # file missing
    p.image_path = "../server.py"
    assert server._product_dict(p)["image_src"] is None  # outside downloads/images
    p.image_path = None
    assert server._product_dict(p)["image_src"] is None


def test_port_override_env():
    assert server._port({"FRIDGE_PORT": "9123"}) == 9123
    assert server._port({}) == 8765 and server._port({"FRIDGE_PORT": "abc"}) == 8765 and server._port({"FRIDGE_PORT": "80"}) == 8765



def _lg_items():
    """Distinct, scrapable LG mock products as collect items (fridges first)."""
    return [{"brand": "LG", "url": u, "subcategory": sub, "model_number": m}
            for m, n, p, major, sub, u in server._MOCK_ROWS["LG"] if m != "LRFOS3016S"]


def _lg_item(i):
    return _lg_items()[i]


def test_max_collect_cap_and_multiple_models_per_brand():
    cap = server.service.MAX_COLLECT
    meta = client.get("/api/meta").json()
    assert cap == 12 and meta["max_collect"] == cap == meta["max_selected"]
    over = _lg_items()[:cap + 1]
    assert len(over) == cap + 1
    assert client.post("/api/collect", json={"urls": over}).status_code == 422
    ok = client.post("/api/collect", json={"urls": over[:cap]})
    assert ok.status_code == 200, ok.text
    wait(ok.json()["job_id"], timeout=120)
    two = client.post("/api/collect", json={"urls": [_lg_item(1), _lg_item(2)]})  # two LG fridges
    assert two.status_code == 200
    res = wait(two.json()["job_id"])["result"]
    assert len(res["groups"]) == 1 and len(res["groups"][0]["products"]) == 2  # brand x group x market, no per-brand cap


def test_regions_endpoint():
    regs = client.get("/api/regions").json()
    assert [r["key"] for r in regs] == ["kr", "na", "eu", "sa", "me", "as", "oc"]
    by = {r["key"]: r for r in regs}
    assert by["na"]["enabled"] and by["na"]["default"] and by["na"]["note"] == "" and by["na"]["currency"] == "USD"
    assert "Samsung" in by["na"]["brands"] and by["na"]["countries"][:2] == ["us", "ca"]
    assert by["na"]["enabled_countries"] == ["us"]
    for k in ("kr", "eu", "sa"):
        assert not by[k]["enabled"] and by[k]["note"] == "준비 중" and by[k]["brands"] == [] and not by[k]["default"]
    assert by["kr"]["countries"] == ["kr"] and by["eu"]["countries"][:2] == ["de", "uk"] and by["kr"]["label_ko"] == "한국"
    assert [r["key"] for r in regs if r["default"]] == ["na"]


def test_filters_endpoint():
    r = client.get("/api/filters", params={"subcategory": "french_door", "region": "na"})
    assert r.status_code == 200, r.text
    body = r.json()
    keys = [g["key"] for g in body["groups"]]
    assert body["subcategory"] == "french_door" and body["region"] == "na" and keys[:2] == ["brand", "price"]
    cap = next(g for g in body["groups"] if g["key"] == "capacity_l")
    assert cap["display_units"] == ["L", "cu ft"] and cap["level"] == "listing" and cap["values"][0]["key"] == "lt400"
    assert {"key", "label_ko", "type", "level", "values", "more_after", "default_open"} <= set(cap)
    assert "door_type" not in keys and "kr_grade" not in keys
    major = client.get("/api/filters", params={"subcategory": "refrigerator"}).json()
    assert "door_type" in [g["key"] for g in major["groups"]] and major["region"] == "na"
    kr = client.get("/api/filters", params={"subcategory": "french_door", "region": "na,kr"}).json()
    assert "kr_grade" in [g["key"] for g in kr["groups"]] and kr["region"] == "na,kr"
    for bad in ({"subcategory": "nope"}, {"subcategory": "french_door", "region": "mars"}, {}):
        assert client.get("/api/filters", params=bad).status_code == 422, bad


def test_categories_and_brands_accept_region():
    na = {c["key"]: c for c in client.get("/api/categories", params={"region": "na"}).json()}
    assert na["refrigerator"]["enabled"]
    kr = client.get("/api/categories", params={"region": "kr"}).json()
    assert all(not c["enabled"] and c["brands"] == [] for m in kr for c in m["children"])
    assert client.get("/api/categories", params={"region": "mars"}).status_code == 422
    brands = {b["name"]: b for b in client.get("/api/brands", params={"region": "kr"}).json()}
    assert not brands["Samsung"]["enabled"] and brands["Samsung"]["categories"] == []
    assert {b["name"]: b for b in client.get("/api/brands").json()}["Samsung"]["regions"] == ["na"]
    assert client.get("/api/brands", params={"region": "x"}).status_code == 422


def test_search_regions_default_and_validation():
    j = run_search(brands=["LG"], category=None, subcategories=["french_door"])
    res = j["result"]
    assert res["regions"] == ["na"] and res["groups"][0]["region"] == "na" and res["groups"][0]["currency"] == "USD"
    c = next(c for k in res["bands"] for c in res["bands"][k])
    assert (c["region"], c["country"], c["currency"], c["price_local"]) == ("na", "us", "USD", None)
    assert isinstance(c["attrs"], dict) and isinstance(c["attrs_src"], dict)
    assert res["filtered_total"] == res["unfiltered_total"] == res["total"] and res["region_totals"] == {"na": res["total"]}
    assert "brand" in res["facets"] and "capacity_l" in res["facets"] and "__unknown__" in res["facets"]["brand"]
    base = {"brands": ["LG"], "subcategories": ["french_door"]}
    for bad in ({"regions": ["kr"]}, {"regions": ["mars"]}, {"regions": []}, {"regions": ["na"] * 5 + ["eu"]},
                {"regions": ["na", "kr"]}):  # disabled / unknown / empty / too many / one disabled region among enabled
        assert client.post("/api/search", json={**base, **bad}).status_code == 422, bad


def test_search_filters_facets_and_validation():
    j = run_search(brands=["Samsung", "LG"], category=None, subcategories=["french_door", "side_by_side"])
    allc = [c for g in j["result"]["groups"] for k in g["bands"] for c in g["bands"][k]]
    lg = [c for c in allc if c["brand"] == "LG"]
    jf = run_search(brands=["Samsung", "LG"], category=None, subcategories=["french_door", "side_by_side"],
                    filters={"brand": ["LG"]})
    res = jf["result"]
    got = [c for g in res["groups"] for k in g["bands"] for c in g["bands"][k]]
    assert {c["brand"] for c in got} == {"LG"} and len(got) == len(lg)
    assert res["filtered_total"] == len(lg) == res["total"] and res["unfiltered_total"] == len(allc)
    f = res["facets"]["brand"]  # disjunctive: the brand facet ignores the brand selection itself
    assert f["Samsung"] == len([c for c in allc if c["brand"] == "Samsung"]) and f["LG"] == len(lg)
    assert sum(res["facets"]["capacity_l"].values()) == len(lg)  # every LG candidate in exactly one bucket (+unknown)
    big = run_search(brands=["LG"], category=None, subcategories=["french_door"], filters={"capacity_l": ["700_plus"]})
    assert big["result"]["filtered_total"] < big["result"]["unfiltered_total"]
    assert all("cu. ft." in c["name"] for g in big["result"]["groups"] for k in g["bands"] for c in g["bands"][k])
    unk = run_search(brands=["Samsung"], category=None, subcategories=["french_door"],
                     filters={"energy": ["energy_star", "__unknown__"]})
    assert unk["result"]["filtered_total"] == unk["result"]["unfiltered_total"]  # mock names carry no ENERGY STAR
    none = run_search(brands=["Samsung"], category=None, subcategories=["french_door"], filters={"energy": ["energy_star"]})
    assert none["result"]["filtered_total"] == 0 and none["result"]["unfiltered_total"] > 0
    base = {"brands": ["LG"], "subcategories": ["french_door"]}
    for bad in ({"nope": ["x"]}, {"brand": "LG"}, {"capacity_l": ["bogus"]}, {"price": {"min": -1}},
                {"capacity_kg": ["lt9"]}, {"brand": ["x" * 65]}, ["brand"], {f"k{i}": [] for i in range(21)}):
        assert client.post("/api/search", json={**base, "filters": bad}).status_code == 422, bad


def test_search_without_filters_keeps_legacy_shape():
    j = run_search(brands=["LG"], category=None, subcategories=["front_load"])
    res = j["result"]
    assert {"groups", "bands", "labels", "thresholds", "total", "failed_brands"} <= set(res)
    assert set(res["groups"][0]) >= {"category", "label_ko", "bands", "labels", "thresholds", "total"}


def test_collect_items_carry_region_country():
    item = {**_lg_item(5), "country": "us", "region": "na"}
    r = client.post("/api/collect", json={"urls": [item]})
    assert r.status_code == 200, r.text
    res = wait(r.json()["job_id"])["result"]
    assert res["products"] and all(p["region"] == "na" and p["country"] == "us" and p["currency"] == "USD"
                                   for p in res["products"])
    for bad in ({"country": "kr"}, {"country": "zz"}, {"region": "kr"}, {"country": "us", "region": "eu"}, {"region": "mars"}):
        assert client.post("/api/collect", json={"urls": [{**_lg_item(6), **bad}]}).status_code == 422, bad


def test_host_allowlist_extra_domains():
    assert server.host_allowed("LG", "https://www.lge.co.kr/refrigerators/x")
    assert not server.host_allowed("Samsung", "https://www.lge.co.kr/x")
    assert not server.host_allowed("LG", "http://www.lge.co.kr/x")
    assert not server.host_allowed("LG", "https://lge.co.kr.evil.io/x")


def test_match_and_launches_endpoints_on_a_seeded_history():
    from catalog import Candidate
    from store import Store
    prev = (server._STORE, server._SEEN_READY)
    with tempfile.TemporaryDirectory() as d:
        st = Store(Path(d) / "m.db")
        cands = [Candidate(brand=b, model_number=f"{b[:2]}{i:02d}", name=f"30 in. Smart Wall Oven {b}{i}", url=f"https://x.test/{b}/{i}",
                           price_usd=900.0 + 100 * i, category="cooking", subcategory="electric_oven",
                           attrs={"rating": 4.5, "review_count": 80 + i, "release_date": "2026-03"} if i == 3 else {})
                 for b in ("GE", "LG") for i in range(8)]
        st.record_seen(cands, limit=30)
        server._STORE, server._SEEN_READY = st, True
        try:
            r = client.post("/api/match", json={"sub": "electric_oven", "price": 1300, "country": "us",
                                                "specs": {"width_in": 30, "features": ["wifi"]}, "tier_window": 1})
            assert r.status_code == 200, r.text
            out = r.json()
            assert len(out["tiers"]) == 5 and out["target"]["tier"] in (1, 2, 3, 4, 5) and out["results"]
            assert out["data"]["models"] == 16 and out["data"]["with_rating"] == 2 and out["data"]["with_release_date"] == 2
            assert all(abs(x["tier_diff"]) <= 1 for x in out["results"] if x["tier_diff"] is not None)
            top = out["results"][0]
            assert {"price", "spec", "recency", "response"} <= set(top["components"]) and 0 <= top["coverage"] <= 1
            assert client.post("/api/match", json={"sub": "nope"}).status_code == 422
            assert client.post("/api/match", json={"sub": "electric_oven", "country": "xx"}).status_code == 422
            assert client.post("/api/match", json={"sub": "electric_oven", "weights": {"x": 1}}).status_code == 422
            assert client.post("/api/match", json={"sub": "electric_oven", "weights": {"price": 500}}).status_code == 422
            assert client.post("/api/match", json={"sub": "electric_oven", "specs": {"features": ["laser"]}}).status_code == 422
            assert client.post("/api/match", json={"sub": "electric_oven", "band_pct": 1}).status_code == 422
            la = client.get("/api/launches", params={"sub": "electric_oven", "window": 12}).json()
            assert la["totals"]["new"] == 2 and set(la["by_brand"]) == {"GE", "LG"} and la["note"]
            assert client.get("/api/launches", params={"sub": "nope"}).status_code == 422
            assert client.get("/api/launches", params={"sub": "electric_oven", "window": 0}).status_code == 422
        finally:
            server._STORE, server._SEEN_READY = prev


def test_schedule_endpoints_validate_run_conflict_and_delete():
    from catalog import Candidate
    from store import Store
    prev = (server._STORE, os.environ.get("FRIDGE_MOCK"), server.service.search, os.environ.get("FRIDGE_BROWSER_MODE"),
            os.environ.get("FRIDGE_HEADLESS"))
    calls = []

    def fake_search(brands, subs, limit, store=None, use_cache=True, countries=None):
        calls.append((tuple(brands), use_cache, os.environ.get("FRIDGE_BROWSER_MODE")))
        found = [Candidate(brand=brands[0], model_number=f"{brands[0][:2]}{i}", name="Oven", url=f"https://x.test/{brands[0]}/{i}",
                           price_usd=1000.0 + i, category="cooking", subcategory=subs[0], country="us") for i in range(3)]
        store.record_seen(found, limit=limit)
        return found, [(brands[0], "ok", "3 candidates")]

    with tempfile.TemporaryDirectory() as d:
        server._STORE = Store(Path(d) / "sch.db")
        server.service.search = fake_search
        os.environ["FRIDGE_MOCK"] = "0"
        try:
            body = {"name": "주간 오븐", "brands": ["Samsung", "LG"], "subcategories": ["electric_oven"], "regions": ["na"],
                    "limit": 30, "interval_h": 24}
            r = client.post("/api/schedules", json=body)
            assert r.status_code == 200, r.text
            sch = r.json()
            assert sch["enabled"] and sch["combos"] == 2 and sch["running"] is False and sch["last_run"] is None
            assert client.get("/api/schedules").json()["schedules"][0]["id"] == sch["id"]
            assert client.get("/api/schedules").json()["limits"]["min_interval_h"] == 6.0
            for bad in ({**body, "name": ""}, {**body, "brands": ["Nope"]}, {**body, "subcategories": ["nope"]},
                        {**body, "interval_h": 1}, {**body, "limit": 1}, {**body, "regions": ["zz"]},
                        {**body, "brands": ["Amana"], "subcategories": ["sco"]}):
                assert client.post("/api/schedules", json=bad).status_code == 422, bad
            assert client.post("/api/schedules/zzzz/run").status_code == 404
            assert client.post("/api/schedules/000000000000/delete").status_code == 404
            # a busy job slot (the user's own search) refuses a scheduled run
            blocker = server.Job("search")
            blocker.status = "running"
            server.JOBS[blocker.id] = blocker
            assert client.post(f"/api/schedules/{sch['id']}/run").status_code == 409
            server.JOBS.pop(blocker.id)
            job_id = client.post(f"/api/schedules/{sch['id']}/run").json()["job_id"]
            done = wait(job_id)
            assert done["status"] == "done" and done["result"]["summary"]["status"] == "ok"
            assert done["result"]["summary"]["candidates"] == 6 and all(c[1] is False and c[2] == "headless" for c in calls)
            after = client.get("/api/schedules").json()["schedules"][0]
            assert after["last_status"] == "ok" and after["next_run"] > after["last_run"] and after["last_summary"]["brands"] == 2
            # due logic: a disabled schedule never starts, an enabled and overdue one does
            assert client.post(f"/api/schedules/{sch['id']}/toggle", json={"enabled": False}).json()["enabled"] is False
            server._store().update_schedule(sch["id"], next_run="2000-01-01T00:00:00")
            n = len(calls)
            server._scheduler_tick()
            assert len(calls) == n
            client.post(f"/api/schedules/{sch['id']}/toggle", json={"enabled": True})
            server._scheduler_tick()
            end = time.time() + 20
            while len(calls) < n + 2 and time.time() < end:
                time.sleep(0.05)
            assert len(calls) == n + 2
            assert client.post(f"/api/schedules/{sch['id']}/delete").json() == {"ok": True}
            assert client.get("/api/schedules").json()["schedules"] == []
        finally:
            time.sleep(0.3)
            server._STORE, server.service.search = prev[0], prev[2]
            for key, old in (("FRIDGE_MOCK", prev[1]), ("FRIDGE_BROWSER_MODE", prev[3]), ("FRIDGE_HEADLESS", prev[4])):
                if old is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = old


def test_schedule_run_is_refused_in_mock_mode():
    from store import Store
    prev = server._STORE
    with tempfile.TemporaryDirectory() as d:
        server._STORE = Store(Path(d) / "mock.db")  # never touch the real cache in a test
        try:
            r = client.post("/api/schedules", json={"name": "x", "brands": ["Samsung"], "subcategories": ["electric_oven"]})
            assert r.status_code == 200, r.text  # creating is fine in mock mode (it is only stored)...
            assert client.post(f"/api/schedules/{r.json()['id']}/run").status_code == 422  # ...a real re-search is not run
            assert server._runner_enabled() is False
        finally:
            server._STORE = prev


def test_match_page_is_served():
    r = client.get("/match")
    assert r.status_code in (200, 404)  # 404 until web/match.html exists; never a 500
    if r.status_code == 200:
        assert "경쟁 모델" in r.text


if __name__ == "__main__":
    fns =[v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print("ok  ", fn.__name__)
    print(f"{len(fns)} passed")
